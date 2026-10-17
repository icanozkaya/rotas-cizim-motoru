# -*- coding: utf-8 -*-
"""Kullanıcının ISI_KAZANCI_VE_KLIMA_SANTRALI_HESABI_1.xlsx dosyasına KAPAK (proje bilgileri) sayfası ekler — rakamlara dokunmaz.
Kapak bilgisi: SASKİ yazısı 01.09.2026 E-98014594-220.03.03-142845 (KYR ADA YATIRIM, Karakamış Mah., 4265 ada 1 parsel)."""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins
wb = openpyxl.load_workbook('GIRDI_ISI_KAZANCI_1.xlsx')
ws = wb.create_sheet('KAPAK', 0)
F = lambda sz=11, b=False, c='000000': Font(name='Arial', size=sz, bold=b, color=c)
ince = Side(style='thin'); kalin = Side(style='medium')
for col, w in zip('ABCD', (4, 38, 46, 4)):
    ws.column_dimensions[col].width = w
r = 2


def satir(a, b, kal=False, h=22):
    global r
    ws.cell(r, 2, a).font = F(10, True); ws.cell(r, 3, b).font = F(10)
    for c in (2, 3):
        x = ws.cell(r, c); x.border = Border(left=ince, right=ince, top=ince, bottom=ince)
        x.alignment = Alignment(vertical='center', wrap_text=True, horizontal='left', indent=1)
    ws.row_dimensions[r].height = h; r += 1


def baslik(s, h=24, fill='D9D9D9'):
    global r
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ws.cell(r, 2, s); c.font = F(11, True); c.alignment = Alignment(horizontal='center', vertical='center')
    c.fill = PatternFill('solid', fgColor=fill)
    for cc in (2, 3):
        ws.cell(r, cc).border = Border(left=ince, right=ince, top=ince, bottom=ince)
    ws.row_dimensions[r].height = h; r += 1


ws.merge_cells('B2:C2'); ws['B2'] = 'ROTAŞ MÜHENDİSLİK . PROJE . MÜŞAVİRLİK'; ws['B2'].font = F(14, True)
ws['B2'].alignment = Alignment(horizontal='center'); ws.row_dimensions[2].height = 26
ws.merge_cells('B3:C3'); ws['B3'] = 'Semerciler Mah. Sait Faik Sk. No:22/C Z01 Adapazarı/Sakarya  •  0264 272 00 49  •  rotasmuhendislik@gmail.com'
ws['B3'].font = F(8); ws['B3'].alignment = Alignment(horizontal='center')
r = 6
ws.merge_cells(start_row=r, start_column=2, end_row=r + 2, end_column=3)
c = ws.cell(r, 2, 'ISITMA-SOĞUTMA VE HAVALANDIRMA TESİSATI\nISI KAZANCI (m³ YÖNTEMİ) VE ISI GERİ KAZANIMLI\nKLİMA SANTRALİ HESAP RAPORU')
c.font = F(15, True); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for rr in range(r, r + 3):
    ws.row_dimensions[rr].height = 26
    for cc in (2, 3):
        ws.cell(rr, cc).border = Border(left=kalin if cc == 2 else None, right=kalin if cc == 3 else None,
                                         top=kalin if rr == r else None, bottom=kalin if rr == r + 2 else None)
r = 11
baslik('BİNANIN')
satir('Yapı Sahibi', 'KYR ADA YATIRIM İNŞ. SAN. VE TİC. A.Ş.')
satir('Yetkili', 'Oğuzhan ARKIN')
satir('Kullanma Amacı', 'TİCARİ — MARKET VE İŞYERLERİ')
satir('Proje Konusu', 'Isı geri kazanımlı klima santralleri (DX ısıtma / soğutma bataryalı), ısı kazancı ve havalandırma', h=34)
r += 1
baslik('ARSANIN')
satir('İli', 'SAKARYA')
satir('İlçesi', 'ADAPAZARI')
satir('Mahallesi', 'KARAKAMIŞ')
satir('Pafta', '-----')
satir('Ada', '4265')
satir('Parsel', '1')
r += 1
baslik('HESABI YAPAN')
satir('Adı Soyadı', 'BAYRAM GENÇ')
satir('Ünvanı', 'MAKİNA MÜHENDİSİ')
satir('Oda Sicil No', '56858')
satir('Büro', 'ROTAŞ MÜHENDİSLİK')
satir('İmza', '', h=44)
r += 1
baslik('ONAY')
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
for cc in (2, 3):
    ws.cell(r, cc).border = Border(left=ince, right=ince, top=ince, bottom=ince)
ws.row_dimensions[r].height = 70; r += 1
r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
c = ws.cell(r, 2, 'İÇİNDEKİLER: Varsayımlar • Mahaller • KS Hesabı • Fan Hesabı • Cihaz Listesi • Elektrik Yükleri • Kontrol')
c.font = F(9); c.alignment = Alignment(horizontal='center', wrap_text=True); ws.row_dimensions[r].height = 28
r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
c = ws.cell(r, 2, 'BU PROJE MÜELLİFİN İZNİ OLMADAN YAYINLANAMAZ-ÇOĞALTILAMAZ-KULLANILAMAZ-DEĞİŞTİRİLEMEZ.')
c.font = F(8, True); c.alignment = Alignment(horizontal='center', wrap_text=True)
ws.page_setup.orientation = 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.6)
ws.print_options.horizontalCentered = True
ws.sheet_view.showGridLines = False
ws.print_area = f'A1:D{r}'
# diğer sayfalar: değerlere dokunmadan yazdırma ayarı (A4, genişliğe sığdır, uzun sayfalar alta taşar), alt bilgi
for s in wb.worksheets[1:]:
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.oddFooter.left.text = 'ROTAŞ MÜHENDİSLİK — KYR ADA YATIRIM / 4265 ada 1 parsel'
    s.oddFooter.left.size = 8
    s.oddFooter.right.text = '&A — Sayfa &P / &N'; s.oddFooter.right.size = 8
wb.active = 0
wb.save('ISI_KAZANCI_VE_KLIMA_SANTRALI_HESABI_KAPAKLI.xlsx')
print('ok')
# MAHALLER (25 sütun): iki sayfa genişliğinde, A–C sütunları her sayfada tekrar
wb = openpyxl.load_workbook('ISI_KAZANCI_VE_KLIMA_SANTRALI_HESABI_KAPAKLI.xlsx')
m = wb['MAHALLER']
m.page_setup.fitToWidth = 2; m.page_setup.fitToHeight = 1
m.print_title_cols = 'A:C'
wb.save('ISI_KAZANCI_VE_KLIMA_SANTRALI_HESABI_KAPAKLI.xlsx')
