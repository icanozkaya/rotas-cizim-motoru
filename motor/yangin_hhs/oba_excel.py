# -*- coding: utf-8 -*-
"""OBA 84 parsel yangın + duman tahliye hesap Excel'i."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import oba_veri as V
import oba_duman as OD
from ciz_ortak import KOK

B = Border(*(Side(style='thin', color='999999'),) * 4)
BAS = PatternFill('solid', fgColor='DDE6F0')


def sayfa(wb, ad, baslik, kolon, satirlar, gen):
    ws = wb.create_sheet(ad)
    ws['A1'] = baslik
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = 'OBA MAKARNACILIK SAN. ve TİC. A.Ş. - Hendek Çağlayan Mah. 900 ada 84 parsel - Depolama binası | Hazırlayan: Can ÖZKAYA (Teknik Ressam) - ROTAŞ Mühendislik'
    ws['A2'].font = Font(italic=True, size=9)
    for j, k in enumerate(kolon):
        c = ws.cell(4, j + 1, k)
        c.font = Font(bold=True); c.fill = BAS; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, r in enumerate(satirlar):
        for j, v in enumerate(r):
            c = ws.cell(5 + i, j + 1, v)
            c.border = B
            if isinstance(v, float):
                c.number_format = '0.000' if abs(v) < 10 else '0.0'
    for j, g in enumerate(gen):
        ws.column_dimensions[chr(65 + j)].width = g
    ws.freeze_panes = 'A5'
    return ws


def main():
    O = V.ozet()
    S, H, D = O['S'], O['H'], O['D']
    T = V.TASARIM
    wb = Workbook()
    wb.remove(wb.active)
    sayfa(wb, 'Özet', 'YANGIN TESİSATI HESAP ÖZETİ (TS EN 12845 HHS / BYKHY)', ['Büyüklük', 'Değer', 'Birim', 'Açıklama'], [
        ['Tehlike sınıfı', 'HHS', '', 'Depolama yüksek tehlike, ST1 paletli blok istif ≤ 4,0 m, ürün kat. II'],
        ['Tasarım yoğunluğu', T['yogunluk'], 'mm/dk', 'Ç.4: 7,5 → güvenli taraf 10,0'],
        ['İşletme alanı', T['alan'], 'm²', ''],
        ['Süre', T['sure'], 'dk', ''],
        ['Başlık K / Pmin', f"K{T['K']} / {T['Pmin']}", 'bar', 'dik, 68 °C'],
        ['Toplam başlık', len(V.BR_X) * len(V.BAS_Y), 'adet', '2 zon x 576'],
        ['Uzak alan başlık sayısı', S['nA'], 'adet', f"{S['alan']:.0f} m²"],
        ['En düşük başlık debisi', round(S['q_min'], 1), 'l/dk', ''],
        ['Gerçekleşen yoğunluk', round(S['yog'], 2), 'mm/dk', '≥ 10'],
        ['Sprinkler debisi', round(S['Qs'], 0), 'l/dk', ''],
        ['Pompa çıkışı gerekli basınç (sprinkler)', round(S['P0'], 3), 'bar', ''],
        ['En yakın alan debisi', round(S['Qf'], 0), 'l/dk', 'aynı basınçta'],
        ['Hidrant debisi', H['Q'], 'l/dk', '7 bar çıkış'],
        ['Pompa çıkışı gerekli basınç (hidrant)', round(H['P_pompa'], 3), 'bar', ''],
        ['Dolap debisi', D['Q'], 'l/dk', '2 x 100 l/dk, 4 bar'],
        ['Toplam debi', round(O['Q_top'], 0), 'l/dk', f"{O['Qm3']:.1f} m³/h"],
        ['Depo hacmi (90 dk)', round(O['V'], 1), 'm³', f"seçilen {O['V_sec']} m³, 2 göz"],
        ['Dolum debisi (36 sa)', round(O['dolum'], 2), 'm³/h', ''],
        ['Ana pompa Q / H', f"{O['Q_sec']} / {O['H_sec']}", 'm³/h / mSS', 'elektrikli + dizel, her biri %100'],
        ['Mil gücü / motor', f"{O['P_mil']:.1f} / {O['motor']} (dizel {O['dizel']})", 'kW', 'η = 0,74'],
        ['Jokey', f"{O['jokey']['Q']} m³/h / {O['jokey']['H']} mSS / {O['jokey']['P']} kW", '', ''],
    ], [42, 22, 14, 60])
    rows = [[t['u'], t['a'], f"DN{t['dn']}", round(t['L'], 2), t['nd'], t['nt'], round(t['Q'], 1), round(t['v'], 3),
             round(t['j'], 5), round(t['z'], 3), t['ek'], round(t['Pu'], 4), round(t['Pa'], 4)] for t in S['tablo']]
    sayfa(wb, 'Sprinkler hat hesabı', 'HİDROLİK EN UZAK ALAN - HAT HESABI (Hazen-Williams C=120)',
          ['Üst düğüm', 'Alt düğüm', 'Çap', 'L (m)', 'Dirsek', 'Te (dönüş)', 'Q (l/dk)', 'v (m/s)', 'J (bar/m)', 'Δz (m)',
           'Ek kayıp (bar)', 'P üst (bar)', 'P alt (bar)'], rows, [12, 12, 8, 8, 8, 9, 10, 9, 11, 8, 11, 11, 11])
    bas = [[h, round(S['yol'][h][0]), round(S['yol'][h][1]), round(S['P'][h], 4), round(S['Qb'][h], 2)] for h in S['uzak']]
    sayfa(wb, 'Uzak alan başlıklar', 'HİDROLİK EN UZAK ALAN BAŞLIKLARI (K160)', ['Başlık', 'x (cm)', 'y (cm)', 'P (bar)', 'q (l/dk)'],
          bas, [14, 12, 12, 12, 12])
    sayfa(wb, 'Hidrant-Dolap', 'HİDRANT VE YANGIN DOLABI HESABI', ['Büyüklük', 'Değer', 'Birim'], [
        ['Hidrant adedi', H['n'], 'adet'], ['Halka uzunluğu', round(H['L_halka'], 1), 'm'], ['Halka çapı', 'DN150 (Ø160 PE100)', ''],
        ['Hidrant debisi', H['Q'], 'l/dk'], ['En uzak hidrant yol boyu', round(H['L_yol'], 1), 'm'],
        ['Halka hızı (Q/2)', round(H['v_halka'], 2), 'm/s'], ['Halka kaybı', round(H['dP_halka'], 3), 'bar'],
        ['Bağlantı + vana kaybı', round(H['dP_bag'], 3), 'bar'], ['Kot farkı', round(H['z'], 2), 'm'],
        ['Pompa çıkışı gerekli', round(H['P_pompa'], 3), 'bar'],
        ['Dolap adedi', len(V.DOLAPLAR), 'adet'], ['Dolap debisi (2 eşzamanlı)', D['Q'], 'l/dk'],
        ['Dolap hat boyu', round(D['L'], 1), 'm'], ['Dolap hattı DN / hız', f"DN{D['dn_hat']} / {D['v']:.2f} m/s", ''],
        ['Dolap kayıp', round(D['dP'], 3), 'bar'], ['Gerekli basınç (halkada)', round(D['P_gerek'], 3), 'bar'],
    ], [40, 24, 10])
    import oba_yangin as OY
    met = []
    OY.METRAJ.clear()
    # metraj geometriden (ozet.json)
    import json
    oz = json.load(open(os.path.join(KOK, 'proje', 'oba_84_parsel_yangin', 'ozet.json')))
    for k, v in oz['metraj'].items():
        met.append([k, round(v, 1), 'm'])
    sayfa(wb, 'Metraj', 'BORU METRAJI (plan geometrisinden)', ['Boru', 'Uzunluk', 'Birim'], met, [40, 14, 8])
    h = OD.hesap()
    sayfa(wb, 'Duman tahliye', 'DUMAN TAHLİYE HESABI (2 seçenek)', ['Büyüklük', 'Değer', 'Birim'], [
        ['Duman bölgesi adedi / alanı', f"8 / {h['A_b']:.0f}", 'm²'], ['Ortalama tavan', round(h['z_ort'], 2), 'm'],
        ['Dumansız yükseklik Y', h['Y'], 'm'], ['Yangın çevresi P', h['P'], 'm'], ['Yangın gücü Q / Qc', f"{h['Qf']:.0f} / {h['Qc']:.0f}", 'kW'],
        ['M = 0,188·P·Y^1,5', round(h['M'], 2), 'kg/s'], ['ΔT', round(h['dT'], 1), 'K'], ['Duman sıcaklığı', round(h['Ts'], 1), '°C'],
        ['Duman yoğunluğu', round(h['rho'], 3), 'kg/m³'], ['Hacimsel debi', round(h['Vs_h'], 0), 'm³/h'],
        ['Tabaka kalınlığı d', round(h['dl'], 2), 'm'],
        ['SEÇENEK 1 (ÇATI) AvCv gerekli', round(h['AvCv'], 2), 'm²/bölge'],
        ['DTK 2,00 x 1,50 Cv 0,60 → Aa', round(h['Aa_dtk'], 2), 'm²/adet'], ['DTK bölge başı (gerekli / seçilen)', f"{h['n_gerek']} / {h['n_dtk']}", 'adet'],
        ['DTK toplam', 8 * h['n_dtk'], 'adet'], ['Taze hava Aa ≥ 1,5·AvCv', round(h['A_taze_dog'], 1), 'm²'],
        ['SEÇENEK 2 (DUVAR) fan', f"{h['n_fan']} x {h['fan_Q']} = {h['Q_fan_top']}", 'm³/h/bölge'],
        ['Fan toplam', 8 * h['n_fan'], 'adet'], ['Fan çıkış hızı', round(h['v_fan'], 1), 'm/s'],
        ['Delinme kontrolü Vmax/fan', round(h['Vmax_fan'], 1), 'm³/s'], ['Taze hava serbest kesit (v ≤ 5 m/s)', round(h['A_taze_mek'], 1), 'm²'],
        ['Taze hava panjur toplam', round(len(OD.PANJUR) * OD.PANJUR_OLCU[0] * OD.PANJUR_OLCU[1] * OD.PANJUR_OLCU[2], 1), 'm²'],
    ], [40, 26, 12])
    yol = os.path.join(KOK, 'proje', 'oba_84_parsel_yangin', 'OBA_84_YANGIN_DUMAN_HESAP.xlsx')
    wb.save(yol)
    print(yol)


if __name__ == '__main__':
    main()
