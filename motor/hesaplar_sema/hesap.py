# -*- coding: utf-8 -*-
"""Şemaya göre HESAPLAR: temiz su kritik devre (YB yöntemi) + blok armatür YB + pis su kolon yükleri (kullanıcı şablonu)"""
import pickle, math, xlrd
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as CL

W0 = xlrd.open_workbook('/tmp/claude-0/hs/ole0.cfb')
S4 = W0.sheet_by_name('Sayfa4')
# Çizelge 3 (plastik): sol tablo YB 0.5-50 sütunlar 15..65 (c=3..9); sağ tablo YB 55-1000 sütunlar 32..100 (c=17..22)
def tablo():
    L, Rt = [], []
    for r in range(5, S4.nrows):
        yb = S4.cell_value(r, 1)
        if isinstance(yb, float):
            L.append((yb, {d: S4.cell_value(r, c) for d, c in zip((15, 20, 25, 32, 40, 50, 65), range(3, 10))}))
        yb2 = S4.cell_value(r, 15)
        if isinstance(yb2, float):
            Rt.append((yb2, {d: S4.cell_value(r, c) for d, c in zip((32, 40, 50, 65, 80, 100), range(17, 23))}))
    return L, Rt
TL, TR = tablo()

def ozgul(yb, dn):
    T = TL if yb <= 50 else TR
    rows = [(y, v.get(dn)) for y, v in T if isinstance(v.get(dn), float)]
    lo = max([r for r in rows if r[0] <= yb], default=rows[0], key=lambda r: r[0])
    hi = min([r for r in rows if r[0] >= yb], default=rows[-1], key=lambda r: r[0])
    if hi[0] == lo[0]: return lo[1]
    return round(lo[1] + (hi[1] - lo[1]) * (yb - lo[0]) / (hi[0] - lo[0]))

PA = 9807
H_HID = 67.10          # şemadaki hidrofor basma yüksekliği (mSS)
KAT_H = 3.00           # kat yüksekliği (m) — varsayım
N_KAT = 7              # bodrum→zemin→1→2→3→4→5→çatı katı döşemesi
Z_ARM = 1.50           # en üst armatür döşemeden yüksekliği
STAT = N_KAT * KAT_H + Z_ARM
AKMA = 49035           # Çizelge 1: musluk DN15 / duş bataryası
SAYAC = 5 * PA         # sayaç kaybı 5 mSS (şablon notu)
# kritik devre: B1 blok çatı katı en uzak armatürden hidrofora (şemadaki YB ve çaplar)
SEG = [
    (1, 'Çatı katı - kritik armatür bağlantısı (şemada 1)', 1.0, 2.5, 20),
    (2, 'Çatı katı daire içi branşman (şemada 2)', 3.0, 3.75, 25),
    (3, '5. kat → çatı katı kolon/branşman (şemada 3, 1 1/4")', 5.0, 24.0, 32),
    (4, 'B1 kolon 4. kat → 5. kat (1 1/2")', KAT_H, 48.0, 40),
    (5, 'B1 kolon 3. kat → 4. kat (2")', KAT_H, 73.0, 50),
    (6, 'B1 kolon 2. kat → 3. kat (2")', KAT_H, 98.0, 50),
    (7, 'B1 kolon 1. kat → 2. kat (2")', KAT_H, 123.0, 50),
    (8, 'B1 kolon zemin → 1. kat (2")', KAT_H, 148.0, 50),
    (9, 'B1 kolon bodrum → zemin (2")', KAT_H, 173.0, 50),
    (10, 'Hidrofor → B1 kolon dibi ana hat (şemada 4, 3")', 70.0, 540.0, 80),
]
ROWS = [(n, ad, L, yb, dn, ozgul(yb, dn)) for n, ad, L, yb, dn in SEG]
LTOT = sum(r[2] for r in ROWS)
H_ETK = round(H_HID * PA - (STAT * PA + AKMA + SAYAC))
R_ORT = H_ETK / LTOT
KAYIP = sum(r[2] * r[5] for r in ROWS)

# ---------------- şemadan sayımlar
inis, K = pickle.load(open('/tmp/claude-0/hs/inis.pkl', 'rb'))
FX = pickle.load(open('/tmp/claude-0/hs/fix.pkl', 'rb'))
def blok(x): return 'B3' if x < 276300 else ('B2' if x < 284800 else 'B1')
KOL = defaultdict(Counter)
for i in inis:
    c = Counter(i['fix'])
    if i['kol'] == 'P9' and i['fl'] == 'BODRUM': c['SUZGEC'] = 0
    for k, v in c.items(): KOL[(i['blok'], i['kol'])][k] += v
SZ = {'B1': 5, 'B2': 7, 'B3': 7}            # bodrum yer süzgeçleri (pis su çukuruna)
ARM = Counter((blok(r['p'][0]), r['nm']) for r in FX if r['nm'] and 'KOLON' in r['nm'])
KOMBI_YB = defaultdict(float)
for r in FX:
    if r['nm'] == 'KOMBI1 KOLON': KOMBI_YB[blok(r['p'][0])] += float(r['v']['SO'].split()[0])

thin = Side(style='thin'); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
HB = Font(bold=True); HF = PatternFill('solid', fgColor='DDEBF7')
def box(ws, r0, c0, r1, c1):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            ws.cell(r, c).border = BR
            ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = openpyxl.Workbook()
# ======================= TEMİZ SU KRİTİK DEVRE
ws = wb.active; ws.title = 'TEMİZ SU KRİTİK DEVRE'
for c, w in zip('ABCDEFGH', (9, 52, 12, 10, 11, 17, 17, 12)): ws.column_dimensions[c].width = w
txt = ['YÜK BİRİMİ YÖNTEMİ (YB) - TEMİZ SU KRİTİK DEVRE BASINÇ KAYBI HESABI (B1-B2-B3 BLOK)',
       'Maksimum su hızı: ana borular 1-3 m/s, daire içi borular 1-2 m/s, sessizlik gereken yerler 0,5-0,7 m/s.',
       'ETKİN BASINÇ = HİDROFOR BASINCI - (STATİK KOT FARKI + AKMA BASINCI + SAYAÇ KAYBI)',
       f'Hidrofor basıncı (şemadaki) = {H_HID:.2f} mSS x 9807 = {round(H_HID * PA)} Pa',
       f'Statik kot farkı = {N_KAT} kat x {KAT_H:.2f} m + {Z_ARM:.2f} m armatür yüksekliği = {STAT:.2f} m  →  {round(STAT * PA)} Pa',
       f'Akma basıncı (kritik armatür, Çizelge 1) = {AKMA} Pa ;  Sayaç kaybı = 5 mSS x 9807 = {SAYAC} Pa',
       f'ETKİN BASINÇ = {H_HID:.2f}x9807-({STAT:.2f}x9807+{AKMA}+{SAYAC}) = {H_ETK} Pa',
       f'ORTALAMA ÖZGÜL BASINÇ KAYBI = Etkin basınç / Kritik devre uzunluğu = {H_ETK}/{LTOT:.0f} = {R_ORT:.0f} Pa/m']
for i, t in enumerate(txt, 1):
    ws.cell(i, 1, t).font = HB if i in (1, 7, 8) else Font()
r0 = 10
hdr = ['BORU NO', 'KRİTİK DEVRE ELEMANI', 'UZUNLUK (m)', 'YB', 'ÇAP (DN)', 'ÖZGÜL BASINÇ KAYBI (Pa/m)', 'TOPLAM BASINÇ KAYBI (Pa)', 'DEBİ (l/s)']
for j, h in enumerate(hdr, 1):
    c = ws.cell(r0, j, h); c.font = HB; c.fill = HF
for k, (n, ad, L, yb, dn, oz) in enumerate(ROWS, r0 + 1):
    for j, v in enumerate((n, ad, L, yb, f'DN{dn}', oz), 1): ws.cell(k, j, v)
    ws.cell(k, 7, f'=C{k}*F{k}'); ws.cell(k, 8, f'=ROUND(0.25*SQRT(D{k}),3)')
rt = r0 + len(ROWS) + 1
ws.cell(rt, 2, 'TOPLAM').font = HB; ws.cell(rt, 3, f'=SUM(C{r0+1}:C{rt-1})'); ws.cell(rt, 6, 'Pa').font = HB
ws.cell(rt, 7, f'=SUM(G{r0+1}:G{rt-1})').font = HB
box(ws, r0, 1, rt, 8)
for k in range(r0 + 1, rt): ws.cell(k, 2).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.cell(rt + 2, 1, 'Toplam basınç kaybı < Etkin basınç ise seçilen çap değerleri UYGUNDUR.')
ws.cell(rt + 3, 1, f'=IF(G{rt}<{H_ETK},TEXT(G{rt},"0")&" < {H_ETK} Pa OLDUĞUNDAN UYGUNDUR.",TEXT(G{rt},"0")&" > {H_ETK} Pa - ÇAPLAR BÜYÜTÜLMELİDİR.")').font = HB
notes = ['Debi = 0,25 x √YB (l/s) ile tespit edilir.',
         'Özgül basınç kaybı Çizelge 3 (plastik boru) tablosundan, YB değerine göre doğrusal enterpolasyonla alınmıştır.',
         'Kritik devre: hidrofordan en uzak ve en yüksek blok olan B1 blok çatı katındaki son armatür.',
         'Boru YB ve çap değerleri şemadaki MTY çap etiketlerinden alınmıştır; kat yüksekliği 3,00 m kabul edilmiştir.',
         'Akma basıncı kritik devredeki en son cihaza göre Çizelge 1’den, sayaç kaybı 5 mSS olarak alınmıştır.']
for i, t in enumerate(notes): ws.cell(rt + 5 + i, 1, '*' + t)

# ======================= BLOK SAYFALARI
NM = {'LAVABO1 KOLON': 'LAVABO', 'EVIYE1 KOLON': 'SİFON,EVİYE', 'KLOZET1 KOLON': 'KLOZET,HELA', 'DUS TEKNESI1 KOLON': 'BANYO KÜVETİ-DUŞ',
      'CAMASIR MAKINASI1 KOLON': 'ÇAMAŞIR,BUL.MAK', 'SUZGEC': 'YERSÜZGECİ'}
PSROWS = [('KLOZET,HELA', 10), ('BANYO KÜVETİ-DUŞ', 7), ('SİFON,EVİYE', 4), ('İKİ GÖZLÜ EVİYE', 8), ('LAVABO', 2),
          ('YERSÜZGECİ', 2), ('ÇAMAŞIR,BUL.MAK', 4), ('PİSUVAR', 2)]
for B in ('B1', 'B2', 'B3'):
    ws = wb.create_sheet(f'{B} BLOK')
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 24
    ws.cell(1, 2, f'{B} BLOK - ARMATÜR YÜK BİRİMLERİ VE PİS SU KOLON YÜKLERİ').font = Font(bold=True, size=12)
    # armatür YB tablosu
    hdr = ['ARMATÜR', 'BW', 'YB', 'ADET', 'TOPLAM YB']
    for j, h in enumerate(hdr, 2): c = ws.cell(3, j, h); c.font = HB; c.fill = HF
    arm = [('Lavabo', 'DN15', 0.75, ARM[(B, 'LAVABO1 KOLON')]), ('Eviye', 'DN15', 1.0, ARM[(B, 'EVIYE1 KOLON')]),
           ('Rezervuar,Pisuvar', 'DN15', 0.5, ARM[(B, 'KLOZET1 KOLON')]), ('Küvet,Duş', 'DN15', 2.5, ARM[(B, 'DUS TEKNESI1 KOLON')]),
           ('Çamaşır,Bulaşık Makinesi', 'DN15', 1.5, ARM[(B, 'CAMASIR MAKINASI1 KOLON')]),
           ('Şofben / Kombi', 'DN15', 1.0, ARM[(B, 'KOMBI1 KOLON')])]
    for k, (a, bw, yb, n) in enumerate(arm, 4):
        for j, v in enumerate((a, bw, yb, n), 2): ws.cell(k, j, v)
        ws.cell(k, 6, f'=D{k}*E{k}')
    if KOMBI_YB[B] != ARM[(B, 'KOMBI1 KOLON')]:          # şemada 1 adet 0,5 YB şofben var
        n1 = ARM[(B, 'KOMBI1 KOLON')]; ws.cell(9, 6, KOMBI_YB[B]); ws.cell(9, 7, f'({n1 - 1}x1 + 1x0,5)')
    ws.cell(10, 2, 'TOPLAM').font = HB; ws.cell(10, 6, '=SUM(F4:F9)').font = HB
    box(ws, 3, 2, 10, 6)
    # çap referansı
    for j, h in enumerate(('YB ARALIĞI', 'ÇAP'), 8): c = ws.cell(3, j, h); c.font = HB; c.fill = HF
    for k, (a, b) in enumerate((('0-3', '1/2"'), ('3-8', '3/4"'), ('8-20', '1"'), ('20-35', '1 1/4"'), ('35-50', '1 1/2"'), ('50 ve üstü', '2"')), 4):
        ws.cell(k, 8, a); ws.cell(k, 9, b)
    box(ws, 3, 8, 9, 9)
    ws.cell(11, 2, f'Blok debisi: Q = 0,25 x √YB =')
    ws.cell(11, 6, '=ROUND(0.25*SQRT(F10),2)'); ws.cell(11, 7, 'l/s')
    # pis su kolon adetleri
    kols = sorted({k for b, k in KOL if b == B}, key=lambda s: int(s[1:])) + ['SZ']
    def cnt(kol, row):
        if kol == 'SZ': return SZ[B] if row == 'YERSÜZGECİ' else 0
        c = KOL[(B, kol)]
        return sum(v for n, v in c.items() if NM.get(n) == row)
    r = 14
    ws.cell(r, 2, 'PİS SU - KOLON NO / ADET').font = HB
    ws.cell(r + 1, 2, 'CİNSİ').font = HB; ws.cell(r + 1, 3, 'YÜK (SB)').font = HB
    for j, k in enumerate(kols): c = ws.cell(r + 1, 4 + j, k); c.font = HB; c.fill = HF
    for i, (row, sb) in enumerate(PSROWS):
        rr = r + 2 + i
        ws.cell(rr, 2, row); ws.cell(rr, 3, sb)
        for j, k in enumerate(kols):
            v = cnt(k, row); ws.cell(rr, 4 + j, v if v else None)
        ws.cell(rr, 4 + len(kols), f'=SUM({CL(4)}{rr}:{CL(3 + len(kols))}{rr})')
    ws.cell(r + 1, 4 + len(kols), 'TOPLAM').font = HB
    box(ws, r + 1, 2, r + 1 + len(PSROWS), 4 + len(kols))
    # yükler
    r2 = r + 4 + len(PSROWS)
    ws.cell(r2, 2, 'PİS SU - KOLON YÜKLERİ (SB)').font = HB
    ws.cell(r2 + 1, 2, 'CİNSİ').font = HB; ws.cell(r2 + 1, 3, 'YÜK (SB)').font = HB
    for j, k in enumerate(kols): c = ws.cell(r2 + 1, 4 + j, k); c.font = HB; c.fill = HF
    for i, (row, sb) in enumerate(PSROWS):
        rr = r2 + 2 + i; src = r + 2 + i
        ws.cell(rr, 2, row); ws.cell(rr, 3, sb)
        for j in range(len(kols)):
            col = CL(4 + j); ws.cell(rr, 4 + j, f'=$C{rr}*N({col}{src})')
    rt2 = r2 + 2 + len(PSROWS)
    ws.cell(rt2, 2, 'TOPLAM').font = HB
    for j in range(len(kols)):
        col = CL(4 + j); ws.cell(rt2, 4 + j, f'=SUM({col}{r2+2}:{col}{rt2-1})').font = HB
    ws.cell(rt2, 4 + len(kols), f'=SUM({CL(4)}{rt2}:{CL(3+len(kols))}{rt2})').font = HB
    ws.cell(rt2 + 1, 2, 'Qww = 0,5 x √ΣSB (l/s)')
    for j in range(len(kols)):
        col = CL(4 + j); ws.cell(rt2 + 1, 4 + j, f'=ROUND(0.5*SQRT({col}{rt2}),2)')
    box(ws, r2 + 1, 2, rt2 + 1, 4 + len(kols))
    # K.B. tablosu (şablon)
    rk = rt2 + 4
    for j, h in enumerate(('BORU ÇAPI', 'K.B.', 'YÜK DEĞERİ'), 3): c = ws.cell(rk, j, h); c.font = HB; c.fill = HF
    for i, row in enumerate(((50, 10, 10), (70, 40, 25), (100, 150, 100), (125, 400, 270), (150, 900, 600), (200, 2200, 1550))):
        for j, v in enumerate(row, 3): ws.cell(rk + 1 + i, j, v)
    box(ws, rk, 3, rk + 6, 5)
    for j in range(len(kols) + 1): ws.column_dimensions[CL(4 + j)].width = 8
    ws.column_dimensions['C'].width = 10

# ======================= ÇİZELGELER (şablondan aynen)
for nm, title in (('Sayfa2', 'ÇİZELGE_1'), ('Sayfa3', 'ÇİZELGE_2'), ('Sayfa4', 'ÇİZELGE_3')):
    sh = W0.sheet_by_name(nm); ws = wb.create_sheet(title)
    for r in range(sh.nrows):
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            if v != '': ws.cell(r + 1, c + 1, v)
wb.save('/tmp/claude-0/hs/out/HESAPLAR_SIHHI.xlsx')
print('H_ETK', H_ETK, 'L', LTOT, 'R_ort', round(R_ORT), 'KAYIP', round(KAYIP))
for r in ROWS: print(r, r[2] * r[5])
print({B: (sum(n * y for (b, nm), n in ARM.items() if b == B and False), ) for B in ('B1',)})
pickle.dump(dict(ROWS=ROWS, H_ETK=H_ETK, LTOT=LTOT, R_ORT=R_ORT, KAYIP=KAYIP, STAT=STAT, KOL=KOL, ARM=ARM, KOMBI_YB=KOMBI_YB, SZ=SZ), open('/tmp/claude-0/hs/res.pkl', 'wb'))
