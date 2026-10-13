"""Genel: örnek 'yerden ısıtma' şablonundan (formüller korunarak) ısı kaybı + yerden ısıtma + klima + yağmur hesabı."""
import copy, math, pickle, datetime, sys
sys.path.insert(0, '/home/claude/nb'); sys.path.insert(0, '/home/claude/nb/p2')
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import gheat as H
import gklima as GK

SRC = '/tmp/claude-0/xl/ornek_yi.xlsx'


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font); dst.border = copy.copy(src.border); dst.fill = copy.copy(src.fill)
        dst.number_format = src.number_format; dst.alignment = copy.copy(src.alignment)
        dst.protection = copy.copy(src.protection)


def clear_rows(ws, r0, r1, c1=17):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= r0 and rng.max_row <= r1:
            ws.unmerge_cells(str(rng))
    for r in range(r0, r1 + 1):
        for c in range(1, c1 + 1):
            ws.cell(r, c).value = None


# ------------------------------------------------------------------ ısı kaybı sayfası
def write_heat_sheet(ws, tpl, katlbl, sheet_no, rooms, hdr_kat):
    """rooms: [(no, ad, ti, A, rows, qs, ze)]"""
    st = {k: [copy.copy(tpl.cell(r, c)) for c in range(1, 18)] for k, r in
          (('hdr', 8), ('el', 9), ('sum', 29), ('qsl', 30), ('qs', 31), ('top', 32))}
    stc = {k: [(tpl.cell(r, c)) for c in range(1, 18)] for k, r in
           (('hdr', 8), ('el', 9), ('sum', 29), ('qsl', 30), ('qs', 31), ('top', 32))}
    clear_rows(ws, 8, max(ws.max_row, 200))
    ws['O1'] = katlbl
    ws['O2'] = sheet_no
    ws['Q1'] = '=kapak!G41'
    r = 8
    totals = {}
    for no, nm, ti, A, rows, qs, ze in rooms:
        def put(rr, kind, vals):
            for c in range(1, 18):
                copy_style(stc[kind][c - 1], ws.cell(rr, c))
            for col, v in vals.items():
                ws[col + str(rr)] = v
            ws.row_dimensions[rr].height = 10.35 if kind == 'hdr' else 9.95
        put(r, 'hdr', {'F': hdr_kat, 'G': '0', 'H': no, 'I': nm, 'K': '%dºC' % ti})
        ws.merge_cells('I%d:J%d' % (r, r))
        r += 1
        first = r
        rowof = {}
        for i, x in enumerate(rows):
            rowof[x[0]] = first + i
        dd_row = None
        for mark, yon, l, h, n, tk, sub in rows:
            uk = H.ukey(mark)
            label = mark.rstrip('0123456789') if uk not in ('DD25', 'DD30') else uk
            dT = ti - (H.T_DIS if tk is None else tk)
            zh = H.ZH.get(yon, 0) if tk is None else 0
            vals = {'A': label, 'B': (yon or ''), 'D': l, 'E': h, 'F': '=IF(D%d=0," ",D%d*E%d)' % (r, r, r), 'G': n,
                    'I': '=IF(F%d=" "," ",F%d*G%d-H%d)' % (r, r, r, r), 'J': H.U[uk], 'K': dT,
                    'L': '=IF(I%d=" "," ",I%d*J%d*K%d)' % (r, r, r, r), 'M': H.ZD, 'N': H.ZW, 'O': zh,
                    'P': '=IF(M%d=" "," ",(1+(M%d+N%d+O%d)/100))' % (r, r, r, r), 'Q': '=IF(L%d=" "," ",L%d*P%d)' % (r, r, r)}
            if uk == 'DD':
                dd_row = r
            if uk == 'DD25' and dd_row is not None:
                vals['D'] = '=D%d' % dd_row
            if sub:
                vals['H'] = '=' + '+'.join('F%d*G%d' % (rowof[s], rowof[s]) for s in sub)
            put(r, 'el', vals)
            r += 1
        last = r - 1
        put(r, 'sum', {'Q': ('=SUM(Q%d:Q%d)' % (first, last)) if last >= first else 0})
        rs = r
        r += 1
        put(r, 'qsl', {'J': 'Qs=', 'K': 'S a', 'L': 'L', 'M': 'R', 'N': 'H', 'O': 'D  T', 'P': 'Ze'})
        r += 1
        if qs:
            lf = '=' + '+'.join('(D%d+E%d)*2*G%d' % (rowof[m], rowof[m], rowof[m]) for m, a, b, nn in qs)
        else:
            lf = 0
        put(r, 'qs', {'A': ' ', 'J': 'Qs  =', 'K': 2, 'L': lf, 'M': 0.9, 'N': 0.67, 'O': ti - H.T_DIS, 'P': ze,
                      'Q': '=K%d*L%d*M%d*N%d*O%d*P%d' % (r, r, r, r, r, r)})
        rq = r
        r += 1
        put(r, 'top', {'M': '   TOPLAM ISI KAYBI =', 'Q': '=Q%d+Q%d' % (rs, rq)})
        ws.merge_cells('E%d:F%d' % (r, r))
        totals[no] = r
        r += 1
    ws.print_area = 'A1:Q%d' % (r - 1)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return totals


def room_entries(R, kind, no):
    c = H.calc_room(R, kind)
    rows, qs = H.room_rows(R, kind)
    return (no, R['name'], R['ti'], R['A'], rows, qs, c['Ze']), c


# ------------------------------------------------------------------ yardımcı biçimler
THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill('solid', fgColor='D9E1F2')


def cell(ws, ref, v, bold=False, fill=False, fmt=None, al='center', size=9, border=True):
    c = ws[ref]
    c.value = v
    c.font = Font(name='Arial', size=size, bold=bold)
    c.alignment = Alignment(horizontal=al, vertical='center', wrap_text=True)
    if border:
        c.border = BOX
    if fill:
        c.fill = HFILL
    if fmt:
        c.number_format = fmt
    return c


def head(ws, title, sub, P, ncol):
    ws.merge_cells('A1:%s1' % CL(ncol)); cell(ws, 'A1', title, True, size=12, border=False)
    ws.merge_cells('A2:%s2' % CL(ncol)); cell(ws, 'A2', sub, size=8, border=False)
    ws.merge_cells('A3:%s3' % CL(ncol)); cell(ws, 'A3', 'PROJE : %s' % P['proje'], size=9, border=False, al='left')
    ws.merge_cells('A4:%s4' % CL(ncol)); cell(ws, 'A4', 'ROTAŞ MÜHENDİSLİK - Mak. Müh. Bayram GENÇ', size=9, border=False,
                                               al='left')


# ------------------------------------------------------------------ klima sayfası
def klima_sheet(wb, KL, P):
    ws = wb.create_sheet('KLİMA HESABI')
    head(ws, 'DUVAR TİPİ SPLİT KLİMA ISI KAZANCI HESABI VE CİHAZ SEÇİMİ',
         'Yöntem: CLTD / SHGF basit yöntem (TS 2164 - ASHRAE esaslı).  Dış %d °C KT / %d °C YT (Sakarya), iç %d °C %%%d' %
         (GK.T_YAZ_DIS, GK.T_YAZ_YAS, GK.T_YAZ_IC, GK.RH_IC), P, 7)
    for i, w in enumerate([24, 8, 12, 14, 14, 14, 30]):
        ws.column_dimensions[CL(i + 1)].width = w
    r = 6
    cell(ws, 'A%d' % r, 'SABİTLER', True, True); ws.merge_cells('A%d:G%d' % (r, r)); r += 1
    consts = [('Dış duvar U (W/m²K)', GK.U_DUVAR), ('Cam U (W/m²K)', GK.U_CAM), ('Cam gölgeleme katsayısı SC', GK.SC_CAM),
              ('Cam soğutma yük faktörü CLF', GK.CLF_CAM), ('Dış kuru termometre (°C)', GK.T_YAZ_DIS),
              ('İç sıcaklık (°C)', GK.T_YAZ_IC), ('Dış mutlak nem (kg/kg)', GK.W_DIS), ('İç mutlak nem (kg/kg)', GK.W_IC),
              ('m² / kişi', GK.KISI_M2), ('Kişi duyulur (W)', GK.Q_KISI_D), ('Kişi gizli (W)', GK.Q_KISI_G),
              ('Aydınlatma (W/m²)', GK.AYD_W), ('Cihaz (W/m²)', GK.CIH_W), ('Taze hava (l/s.kişi)', GK.TAZE_LS),
              ('Emniyet katsayısı', GK.EMNIYET)]
    CREF = {}
    for nm, v in consts:
        ws.merge_cells('A%d:C%d' % (r, r)); cell(ws, 'A%d' % r, nm, al='left'); cell(ws, 'D%d' % r, v)
        CREF[nm] = '$D$%d' % r
        r += 1
    r += 1
    cell(ws, 'A%d' % r, 'YÖNLERE GÖRE CLTD (K) / MAKS. GÜNEŞ KAZANCI (W/m²)', True, True); ws.merge_cells('A%d:G%d' % (r, r))
    r += 1
    cell(ws, 'A%d' % r, 'YÖN', True, True); cell(ws, 'B%d' % r, 'CLTD', True, True); cell(ws, 'C%d' % r, 'SHG', True, True)
    r += 1
    YREF = {}
    for y in ('K', 'D', 'G', 'B'):
        cell(ws, 'A%d' % r, {'K': 'KUZEY', 'D': 'DOĞU', 'G': 'GÜNEY', 'B': 'BATI'}[y]); cell(ws, 'B%d' % r, GK.CLTD[y])
        cell(ws, 'C%d' % r, GK.SHG[y])
        YREF[y] = r
        r += 1
    r += 1
    sec_rows = []
    for K in KL:
        S, k = K['S'], K['kaz']
        ws.merge_cells('A%d:G%d' % (r, r))
        cell(ws, 'A%d' % r, '%s   (A = %.2f m²)' % (S['name'], S['A']), True, True); r += 1
        for i, h in enumerate(['KALEM', 'YÖN', 'ALAN / ADET', 'U / KATSAYI', 'ΔT / CLTD / q', 'Q (W)', 'AÇIKLAMA']):
            cell(ws, '%s%d' % (CL(i + 1), r), h, True, True)
        r += 1
        cell(ws, 'A%d' % r, 'Alan (m²)', al='left'); cell(ws, 'C%d' % r, S['A']); arow = r; r += 1
        cell(ws, 'A%d' % r, 'Kişi sayısı', al='left'); cell(ws, 'C%d' % r, '=ROUNDUP(C%d/%s,0)' % (arow, CREF['m² / kişi']))
        nrow = r; r += 1
        first = r
        H_ = S.get('H', 2.5)
        for yon, l, wins in S['ext']:
            Ac = sum(w * h * n for w, h, n, _ in wins)
            cell(ws, 'A%d' % r, 'Duvar iletim', al='left'); cell(ws, 'B%d' % r, yon)
            cell(ws, 'C%d' % r, '=%s*%s-%s' % (l, H_, Ac), fmt='0.00'); cell(ws, 'D%d' % r, '=%s' % CREF['Dış duvar U (W/m²K)'])
            cell(ws, 'E%d' % r, '=B%d' % YREF[yon] if False else '=$B$%d' % YREF[yon]); cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0')
            cell(ws, 'G%d' % r, 'L=%.2f m, H=%.2f m' % (l, H_), al='left'); r += 1
            if Ac > 0:
                cell(ws, 'A%d' % r, 'Cam iletim', al='left'); cell(ws, 'B%d' % r, yon)
                cell(ws, 'C%d' % r, Ac, fmt='0.00'); cell(ws, 'D%d' % r, '=%s' % CREF['Cam U (W/m²K)'])
                cell(ws, 'E%d' % r, '=%s-%s' % (CREF['Dış kuru termometre (°C)'], CREF['İç sıcaklık (°C)']))
                cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0')
                cell(ws, 'G%d' % r, ' + '.join('%gx%gx%d' % (w, h, n) for w, h, n, _ in wins), al='left'); r += 1
                cell(ws, 'A%d' % r, 'Cam güneş', al='left'); cell(ws, 'B%d' % r, yon)
                cell(ws, 'C%d' % r, Ac, fmt='0.00')
                cell(ws, 'D%d' % r, '=%s*%s' % (CREF['Cam gölgeleme katsayısı SC'], CREF['Cam soğutma yük faktörü CLF']), fmt='0.00')
                cell(ws, 'E%d' % r, '=$C$%d' % YREF[yon]); cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0')
                cell(ws, 'G%d' % r, 'SC x CLF x SHG', al='left'); r += 1
        for nm, C_, D_, E_, note in [
            ('İnsan duyulur', '=C%d' % nrow, 1, '=%s' % CREF['Kişi duyulur (W)'], 'kişi x W'),
            ('Aydınlatma', '=C%d' % arow, 1, '=%s' % CREF['Aydınlatma (W/m²)'], 'm² x W/m²'),
            ('Cihaz / ekipman', '=C%d' % arow, 1, '=%s' % CREF['Cihaz (W/m²)'], 'm² x W/m²'),
            ('Taze hava duyulur', '=C%d*%s/1000' % (nrow, CREF['Taze hava (l/s.kişi)']), 1.2 * 1006,
             '=%s-%s' % (CREF['Dış kuru termometre (°C)'], CREF['İç sıcaklık (°C)']), 'V(m³/s) x ρcp x ΔT')]:
            cell(ws, 'A%d' % r, nm, al='left'); cell(ws, 'C%d' % r, C_, fmt='0.000'); cell(ws, 'D%d' % r, D_)
            cell(ws, 'E%d' % r, E_); cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0'); cell(ws, 'G%d' % r, note, al='left')
            r += 1
        last_d = r - 1
        cell(ws, 'A%d' % r, 'DUYULUR TOPLAM', True); cell(ws, 'F%d' % r, '=SUM(F%d:F%d)' % (first, last_d), True, fmt='0')
        rd = r; r += 1
        g0 = r
        cell(ws, 'A%d' % r, 'İnsan gizli', al='left'); cell(ws, 'C%d' % r, '=C%d' % nrow); cell(ws, 'D%d' % r, 1)
        cell(ws, 'E%d' % r, '=%s' % CREF['Kişi gizli (W)']); cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0'); r += 1
        cell(ws, 'A%d' % r, 'Taze hava gizli', al='left')
        cell(ws, 'C%d' % r, '=C%d*%s/1000' % (nrow, CREF['Taze hava (l/s.kişi)']), fmt='0.000'); cell(ws, 'D%d' % r, 1.2 * 2501e3)
        cell(ws, 'E%d' % r, '=%s-%s' % (CREF['Dış mutlak nem (kg/kg)'], CREF['İç mutlak nem (kg/kg)']), fmt='0.0000')
        cell(ws, 'F%d' % r, '=C%d*D%d*E%d' % (r, r, r), fmt='0'); cell(ws, 'G%d' % r, 'V x ρ x hfg x Δw', al='left'); r += 1
        cell(ws, 'A%d' % r, 'GİZLİ TOPLAM', True); cell(ws, 'F%d' % r, '=SUM(F%d:F%d)' % (g0, r - 1), True, fmt='0')
        rg = r; r += 1
        cell(ws, 'A%d' % r, 'TOPLAM ISI KAZANCI (x emniyet)', True, True)
        cell(ws, 'F%d' % r, '=(F%d+F%d)*%s' % (rd, rg, CREF['Emniyet katsayısı']), True, True, fmt='0'); rt = r
        cell(ws, 'G%d' % r, '=TEXT(F%d/1000,"0.00")&" kW  =  "&TEXT(F%d*3.412,"0")&" BTU/h"' % (r, r), al='left'); r += 1
        cell(ws, 'A%d' % r, 'ISI KAYBI (TS 825, -6 °C)', True); cell(ws, 'F%d' % r, round(K['kay']['Q'], 1), True, fmt='0')
        cell(ws, 'G%d' % r, 'ZEMİN sayfası ile aynı yöntem', al='left'); rk = r; r += 1
        cell(ws, 'A%d' % r, 'SEÇİLEN CİHAZ', True, True)
        cell(ws, 'C%d' % r, K['adet']); cell(ws, 'D%d' % r, 'adet x'); cell(ws, 'E%d' % r, K['btu']); cell(ws, 'F%d' % r, '=C%d*%s*1000' % (r, K['kw']), fmt='0')
        cell(ws, 'G%d' % r, '=IF(AND(F%d>=F%d,F%d*1.1>=F%d),"UYGUN - inverter duvar tipi split (soğutma/ısıtma)","KONTROL")' %
             (r, rt, r, rk), al='left')
        sec_rows.append(r)
        r += 2
    ws.print_area = 'A1:G%d' % r
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ------------------------------------------------------------------ yağmur sayfası
def yagmur_sheet(wb, P, Pm2, f):
    ws = wb.create_sheet('YAĞMUR SUYU')
    head(ws, 'YAĞMUR SUYU İNİŞ BORUSU ÇAP HESABI', 'ROTAŞ yöntemi: S = P x 0,75 ; A = S / f ; D = 1,13 x √A', P, 4)
    for i, w in enumerate([46, 14, 10, 30]):
        ws.column_dimensions[CL(i + 1)].width = w
    rows = [('P = Çatı (+ balkon) yatay izdüşüm alanı', Pm2, 'm²'),
            ('S = P x 0,75 / 1', '=B6*0.75', 'cm²'),
            ('f = Yağmur iniş borusu adedi', f, 'adet'),
            ('A = S / f', '=B7/B8', 'cm²'),
            ('D = 1,13 x √A', '=1.13*SQRT(B9)', 'cm'),
            ('SEÇİLEN BORU', '=IF(B10<=7,"Ø70 YETERLİDİR - PROJEDE Ø100 KULLANILMIŞTIR","Ø100 UYGUNDUR")', '')]
    r = 6
    for a, b, c in rows:
        cell(ws, 'A%d' % r, a, al='left'); cell(ws, 'B%d' % r, b, fmt='0.00'); cell(ws, 'C%d' % r, c)
        r += 1
    cell(ws, 'A%d' % (r + 1), 'Not: Yağmur suları pis su hattına bağlanmayacak, yağmur suyu kanalına verilecektir.', al='left',
         border=False)
    ws.print_area = 'A1:D%d' % (r + 1)
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ------------------------------------------------------------------ devre hidroliği
def hydraulics(circ):
    out = []
    for d in circ:
        m = d['W'] / (1.163 * 10.0)
        q = m / 3600.0 / 990.0
        v = q / (math.pi * 0.006 ** 2)
        Re = v * 0.012 / 0.6e-6
        lam = 0.316 / Re ** 0.25 if Re > 2300 else 64.0 / max(Re, 1)
        Rpm = lam / 0.012 * 990 * v * v / 2
        out.append(dict(d, m=m, v=v, R=Rpm, dP=Rpm * d['L']))
    return out


# ------------------------------------------------------------------ hidrofor + pis su pompası sayfası
def hidrofor_sheet(wb, P, n_daire, yb_daire, z_dep, z_top, sb_tot, z_pit, z_out):
    ws = wb.create_sheet('HİDROFOR-POMPA')
    head(ws, 'HİDROFOR, SU DEPOSU VE PİS SU POMPASI HESABI', 'TS 1258 / TS EN 806-3 / TS EN 12056-4', P, 4)
    for i, w in enumerate([50, 14, 12, 34]):
        ws.column_dimensions[CL(i + 1)].width = w
    rows = [
        ('A. KULLANMA SUYU DEPOSU', None, None, None),
        ('Bağımsız bölüm sayısı', n_daire, 'adet', ''),
        ('Depo hacmi (≤10 BB için 3 m³, her ek BB +0,5 m³)', '=3+MAX(0,B7-10)*0.5', 'm³', 'mimari depo 3,00 m³'),
        ('B. HİDROFOR', None, None, None),
        ('Daire yük birimi', yb_daire, 'YB', 'TS 1258 Çizelge 15'),
        ('Toplam yük birimi ΣYB', '=B7*B10', 'YB', ''),
        ('Pik debi Q = 0,25·√ΣYB', '=0.25*SQRT(B11)', 'l/s', ''),
        ('Pik debi', '=B12*3.6', 'm³/h', ''),
        ('Depo çıkış kotu', z_dep, 'm', '2. bodrum'),
        ('En yüksek musluk kotu', z_top, 'm', '1. kat duş başlığı'),
        ('Geodezik yükseklik Hg', '=B15-B14', 'mSS', ''),
        ('En uç musluk akış basıncı (kombi)', 15, 'mSS', ''),
        ('Boru + sayaç + vana kayıpları', 4, 'mSS', ''),
        ('Pompa basma yüksekliği H', '=B16+B17+B18', 'mSS', ''),
        ('Devreye girme basıncı Pmin', '=B19/10', 'bar', ''),
        ('Devreden çıkma basıncı Pmax', '=B20+1.5', 'bar', ''),
        ('Saatlik devreye girme sayısı z', 20, '1/h', ''),
        ('Membranlı tank hacmi V = Q·1000/(4·z)·(Pmax+1)/(Pmax−Pmin)', '=B13*1000/(4*B22)*(B21+1)/(B21-B20)', 'lt', '200 lt seçilmiştir'),
        ('SEÇİM', '="2 pompalı (1+1) paket hidrofor Q="&TEXT(B13,"0.0")&" m³/h H="&TEXT(B19,"0")&" mSS"', '', ''),
        ('C. PİS SU POMPASI (1. ve 2. bodrum)', None, None, None),
        ('Pompaya gelen toplam su birimi ΣSB', sb_tot, 'SB', 'D-1 + depo odası'),
        ('Debi Q = K·√ΣSB (K=0,5 konut)', '=0.5*SQRT(B26)', 'l/s', 'TS EN 12056-2'),
        ('Debi', '=B27*3.6', 'm³/h', ''),
        ('Çukur taban kotu', z_pit, 'm', ''),
        ('Basma hattı çıkış kotu (R1 girişi)', z_out, 'm', ''),
        ('Statik yükseklik', '=B30-B29', 'mSS', ''),
        ('Boru + çek valf + vana kayıpları', 2.5, 'mSS', ''),
        ('Pompa basma yüksekliği H', '=B31+B32', 'mSS', ''),
        ('Faydalı çukur hacmi (pompa 10 dk debisinin 1/4ü, min 0,3 m³)', '=MAX(0.3,B28*10/60/4)', 'm³', '100x100x120 cm çukur'),
        ('SEÇİM', '="2 adet (1+1) dalgıç pis su pompası Q="&TEXT(B28,"0.0")&" m³/h H="&TEXT(B33,"0.0")&" mSS"', '', ''),
    ]
    r = 6
    for a, b, c, d in rows:
        if b is None:
            ws.merge_cells('A%d:D%d' % (r, r)); cell(ws, 'A%d' % r, a, True, True, al='left')
        else:
            cell(ws, 'A%d' % r, a, al='left'); cell(ws, 'B%d' % r, b, fmt='0.00'); cell(ws, 'C%d' % r, c); cell(ws, 'D%d' % r, d, al='left')
        r += 1
    ws.print_area = 'A1:D%d' % r
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ------------------------------------------------------------------ kombi pompası + genleşme + kangal sayfası
def kombi_sheet(wb, P, rows):
    """rows: [(kat_daire_adı, circ)]"""
    import ghidrolik as G
    ws = wb.create_sheet('KOMBİ-POMPA-GENLEŞME')
    head(ws, 'KOMBİ POMPASI, GENLEŞME TANKI VE PEX KANGAL KONTROLÜ',
         'Pompa: H = Δp(en uzun devre, Churchill, yerel kayıp x1,3) + kollektör 5 kPa ;  '
         'Genleşme: Vmin = (P2+1,013)/(P2-P1) x (0,01 V + ΔV) , 10→50 °C', P, 12)
    hdr = ['DAİRE', 'YÜK (W)', 'DEBİ (l/h)', 'Δp maks (kPa)', 'H (kPa)', 'Kombi pompası (kPa)', 'POMPA',
           'Sistem suyu (l)', 'ΔV (l)', 'P1 / P2 (bar)', 'Vmin (l)', 'TANK (kombi 8 l)']
    for i, (h, w) in enumerate(zip(hdr, [18, 10, 10, 11, 9, 12, 12, 11, 8, 11, 9, 14])):
        ws.column_dimensions[CL(i + 1)].width = w
        cell(ws, '%s6' % CL(i + 1), h, True, True)
    r = 7
    kang = []
    for nm, circ in rows:
        h = G.daire(circ); g = h['gen']
        vals = [nm, round(h['W']), round(h['Q']), round(h['dpmax'], 2), '=D%d+%g' % (r, G.DP_KOLL), G.KOMBI_POMPA_KPA,
                '=IF(E%d<=F%d,"UYGUN","POMPA GRUBU")' % (r, r), round(g['V'], 1), round(g['dV'], 2),
                '%.1f / %.1f' % (g['P1'], g['P2']),
                '=(%.1f+1.013)/(%.1f-%.1f)*(0.01*H%d+I%d)' % (g['P2'], g['P2'], g['P1'], r, r),
                '=IF(K%d<=%g,"YETERLİ","İLAVE TANK")' % (r, G.KOMBI_TANK_L)]
        for i, v in enumerate(vals):
            cell(ws, '%s%d' % (CL(i + 1), r), v, fmt='0.00' if i in (3, 4, 8, 10) else None)
        kang.append((nm, h))
        r += 1
    r += 1
    cell(ws, 'A%d' % r, 'PEX 16x2,0 KANGAL METRAJI (devre ortasında ek yapılmaz)', True, border=False, al='left'); r += 1
    for i, h in enumerate(['DAİRE', 'TOPLAM (m)', '200 m kangal', '500 m kangal']):
        cell(ws, '%s%d' % (CL(i + 1), r), h, True, True)
    r += 1
    for nm, h in kang:
        for i, v in enumerate([nm, round(h['L'], 1), h['kangal200'], h['kangal500']]):
            cell(ws, '%s%d' % (CL(i + 1), r), v)
        r += 1
    notes = ['Kombi pompasının kullanılabilir basma yüksekliği (%g kPa) seçilen kombinin kataloğundan teyit edilecektir.' % G.KOMBI_POMPA_KPA,
             'Kombi içi genleşme tankı %g l, emniyet ventili 3 bar, ilk dolum basıncı 1,0-1,5 bar alınmıştır.' % G.KOMBI_TANK_L,
             'Devre debileri kollektördeki debimetreli reglaj vanaları ile ayarlanacaktır.']
    r += 1
    for n in notes:
        cell(ws, 'A%d' % r, n, al='left', border=False); r += 1
    ws.print_area = 'A1:L%d' % r
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def logo_temizle(wb):
    """şablondan gelen yabancı logo/resimleri tüm sayfalardan kaldır (ROTAŞ dışı logo kullanılmaz)"""
    n = 0
    for ws in wb.worksheets:
        n += len(getattr(ws, '_images', []))
        ws._images = []
    return n
