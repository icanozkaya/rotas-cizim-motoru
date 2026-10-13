# -*- coding: utf-8 -*-
"""VOYAP mekanik hesap Excel: yangın sayfaları (xl.py) + ısı kaybı, ısıtma, sıhhi, pis su, havalandırma."""
import sys, math
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
sys.path.insert(0, '/home/claude/voyap'); sys.path.insert(0, '/home/claude/voyap/tes')
from hesap import hesapla, TD, U, KH
from veri import *
from sihhi import inch, YB_S, YB_H, PIS_KOL

R = hesapla()
b = Side(style='thin'); B = Border(left=b, right=b, top=b, bottom=b)
HF = PatternFill('solid', fgColor='DDDDDD'); BOLD = Font(bold=True)


def head(ws, t2):
    ws['A1'] = 'ROTAŞ MÜHENDİSLİK - VOYAP - MEKANİK TESİSAT HESAPLARI'; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = t2; ws['A2'].font = Font(bold=True, size=10)


def grid(ws, r0, hdr, rows, widths=None):
    for j, h in enumerate(hdr):
        c = ws.cell(r0, j + 1, h); c.font = BOLD; c.fill = HF; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    r = r0 + 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(r, j + 1, v); c.border = B
        r += 1
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = w
    return r


def build(src, out):
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        ws.title = 'YANGIN-' + ws.title[:24] if not ws.title.startswith('YANGIN') else ws.title
    # ---------------- ısı kaybı
    ws = wb.create_sheet('ISI KAYBI')
    head(ws, f'ISI KAYBI HESABI - TS EN 12831 basitleştirilmiş (Td = {TD} °C, Ti = 20 °C, ısıtılmayan ortak alan 10 °C, otopark 5 °C)')
    ws['A3'] = (f"U değerleri (W/m²K, ısı köprüsü dahil): dış duvar {U['duvar']}, pencere {U['pencere']}, çatı {U['cati']}, "
                f"otopark üstü döşeme {U['doseme_oto']}, iç duvar {U['ic_duvar']}.  Yön zammı K +%5 / G -%5, işletme zammı %7.")
    r = 5
    ozet = []
    for key in sorted(R['ofis']) + list(R['dukkan']):
        o = R['ofis'][key] if key in R['ofis'] else R['dukkan'][key]
        ws.cell(r, 1, o['ad']).font = BOLD
        ws.cell(r, 4, f"Net alan {o['A']} m², hacim {o['V']} m³")
        r += 1
        rows = []
        r0 = r + 1
        for i, (el, Lm, A, Uv, dT, q) in enumerate(o['satir']):
            k = q / (A * Uv * (dT if dT != '' else 1)) if A and Uv else 1
            rr = r0 + i
            rows.append((el, Lm if Lm != '' else None, A, Uv, dT if dT != '' else 1, round(k, 3), f"=ROUND(C{rr}*D{rr}*E{rr}*F{rr},0)"))
        r = grid(ws, r, ('Eleman', 'Uzunluk (m)', 'Alan / Hacim', 'U / katsayı', 'ΔT (K)', 'Zam çarpanı', 'Q (W)'), rows)
        ws.cell(r, 6, 'TOPLAM').font = BOLD
        ws.cell(r, 7, f"=SUM(G{r0}:G{r - 1})").font = BOLD
        ws.cell(r, 8, f"=G{r}/{o['A']}"); ws.cell(r, 9, 'W/m²')
        ozet.append((o['ad'], o['A'], f"='ISI KAYBI'!G{r}"))
        r += 2
    for j, w in enumerate((42, 12, 13, 12, 9, 12, 12, 10, 8)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = w
    # ---------------- ısıtma özet
    ws = wb.create_sheet('ISITMA')
    head(ws, 'ISITMA SİSTEMİ: KAZAN, DEVRELER, POMPALAR, GENLEŞME, BACA, YERDEN ISITMA / RADYATÖR')
    rows = []
    for (kat, k), o in sorted(R['ofis'].items()):
        rows.append((o['ad'], o['A'], o['Q'], o['wm2'], 'Yerden ısıtma 45/35 °C, PEX 16x2,0 / 15 cm (MTY devreleri)', 'Ultrasonik kalorimetre DN15, qp 0,6 m³/h', o['yi']))
    for nm, d in R['dukkan'].items():
        rd = d['rad']
        rows.append((d['ad'], d['A'], d['Q'], d['wm2'], f"{rd['adet']} ad. {rd['tip']} / {int(rd['boy'] * 100)} mm ({rd['wm']} W/m, 70/50/20)", 'Ultrasonik kalorimetre DN20, qp 1,5 m³/h', 'UYGUN'))
    r = grid(ws, 4, ('Bağımsız bölüm', 'Alan m²', 'Q (W)', 'W/m²', 'Isıtıcı', 'Isı sayacı', 'Kontrol'), rows, (30, 9, 9, 8, 52, 36, 9))
    k = R['kazan']; bo = R['boyler']; g = R['genlesme']; bc = R['baca']
    rows = [
        ('Ofis ısı kaybı toplamı', 'Qo', R['Q_ofis'], 'kW', ''),
        ('Dükkân ısı kaybı toplamı', 'Qd', R['Q_dukkan'], 'kW', ''),
        ('Boyler serpantin gücü', 'Qb', bo['P_serp'], 'kW', 'Sıhhi sayfası'),
        ('Kazan hesap gücü', 'Qk', f"=ROUND((C{r + 2}+C{r + 3})*1.1+C{r + 4}*0.5,1)", 'kW', '1,10 x ısıtma + 0,5 x boyler (öncelikli değil)'),
        ('Seçilen kazan', '', f"{k['adet']} x {k['birim']} kW = {k['top']} kW", '', k['tip']),
        ('Doğalgaz tüketimi', 'Vg', R['gaz']['Q'], 'm³/h', 'Q / (8250 x 0,98)'),
        ('Baca kesiti (F = 0,012 Q/√h)', 'F', bc['F'], 'cm²', f"h={bc['h']} m → Ø{bc['Dsec']} (TS EN 13384 ile teyit)"),
        ('Kazan dairesi alt menfez', '', R['kd_hava']['alt'], 'cm²', '5 cm²/kW, min 300 (gaz şartnamesiyle kesinleşir)'),
        ('Kazan dairesi üst menfez', '', R['kd_hava']['ust'], 'cm²', 'alt menfezin %50si'),
        ('Tesisat su hacmi', 'Vs', g['Vs'], 'l', 'Y.I. 18,5 l/kW, panel 8,3 l/kW, kazan 3 l/kW'),
        ('Genleşme hacmi (70 °C)', 'ΔV', g['dV'], 'l', 'n = 0,0224'),
        ('Ön basınç', 'Po', g['Po'], 'bar', 'statik yükseklik/10'),
        ('Genleşme tankı', 'Vn', g['Vn'], 'l', f"seçilen {g['Vsec']} L, emniyet 3 bar"),
    ]
    r = grid(ws, r + 1, ('Kalem', 'Sembol', 'Değer', 'Birim', 'Açıklama'), rows) + 1
    rows = []
    for ad, d in R['devre'].items():
        p = R['pompa'][ad]
        rows.append((ad, d['kw'], d['dT'], d['q'], d['dn'], p['H'], p['P']))
    grid(ws, r, ('Devre', 'Q (kW)', 'ΔT (K)', 'Debi (m³/h)', 'Ana boru', 'Pompa H (mSS)', 'Pompa P (kW)'), rows)
    # ---------------- sıhhi
    ws = wb.create_sheet('SIHHI')
    head(ws, 'TEMİZ SU / SICAK SU / SİRKÜLASYON / HİDROFOR / DEPO / BOYLER (TS 1258, kullanıcının MTY yük birimleri)')
    rows = [('Ofis soğuk su YB (lavabo 0,5 + klozet 0,5 + duş 2,5 + eviye 1,0)', YB_S, inch(YB_S)),
            ('Ofis sıcak su YB (lavabo 0,5 + duş 2,5 + eviye 1,0)', YB_H, inch(YB_H)),
            ('Kat başına Ş1 (6 ofis) soğuk', 6 * YB_S, inch(6 * YB_S)), ('Kat başına Ş2 (8 ofis) soğuk', 8 * YB_S, inch(8 * YB_S)),
            ('Ş1 kolon (2 kat) soğuk', 12 * YB_S, inch(12 * YB_S)), ('Ş2 kolon (2 kat) soğuk', 16 * YB_S, inch(16 * YB_S)),
            ('Ş1 kolon sıcak', 12 * YB_H, inch(12 * YB_H)), ('Ş2 kolon sıcak', 16 * YB_H, inch(16 * YB_H)),
            ('Dükkân bekleme ucu (5 ad.)', 5 * 1.0, inch(5.0)), ('Bina toplam soğuk', R['yb']['top'], inch(R['yb']['top']))]
    r = grid(ws, 4, ('Hat', 'YB', 'Çap'), rows, (64, 10, 10, 12, 30))
    h = R['hidrofor']; d = R['depo']; bo = R['boyler']
    rows = [('Eşzamanlı debi Q = 0,25√ΣYB', f"=ROUND(0.25*SQRT({R['yb']['top']})*3.6,2)", 'm³/h'),
            ('Geometrik yükseklik (depo tabanı → en üst armatür)', h['Hgeo'], 'mSS'),
            ('Pompa alt basınç (akma 15 mSS + kayıp)', h['Halt'], 'mSS'), ('Pompa üst basınç (frekans konvertörlü +1 bar)', h['Hust'], 'mSS'),
            ('Pompa gücü (η=0,55)', h['P'], 'kW'), ('Membranlı tank', h['Vn'], 'l'), ('Hidrofor tipi', h['tip'], ''),
            ('Kullanım suyu deposu', f"{d['kisi']} kişi x {d['l']} l + {d['dukkan']} dükkân x 100 l = {d['V']} m³", f"seçilen {d['Vsec']} m³ ({d['adet']} ad.)")]
    r = grid(ws, r + 1, ('Hidrofor / depo', 'Değer', 'Birim'), rows)
    rows = [(n, a, q, a * q) for n, a, q in bo['cihaz']]
    r0 = r + 2
    r = grid(ws, r + 1, ('Cihaz (sıcak su, 60 °C)', 'Adet', 'l/h', 'Toplam l/h'), rows)
    rows = [('Anlık toplam', f"=SUM(D{r0}:D{r - 1})", 'l/h'), ('Kullanma katsayısı (ofis)', bo['phi'], ''),
            ('Saatlik ihtiyaç', f"=ROUND(B{r + 2}*B{r + 3},0)", 'l/h'), ('Depolama katsayısı', bo['dep'], ''),
            ('Boyler hacmi', f"=ROUND(B{r + 4}*B{r + 5},0)", 'l'), ('Seçilen boyler', f"{bo['Vsec']} L çift serpantinli", ''),
            ('Isıtma gücü (60-10 °C)', bo['P'], 'kW'), ('Serpantin (depolama payı x0,75)', bo['P_serp'], 'kW')]
    grid(ws, r + 1, ('Boyler', 'Değer', 'Birim'), rows)
    # ---------------- pis su
    ws = wb.create_sheet('PIS SU')
    head(ws, 'PİS SU KOLONLARI (TS EN 12056-2, K = 0,5 ofis) - kullanıcının MTY bağlantılarındaki SB değerleri')
    rows = []
    for j, pk in enumerate(PIS_KOL):
        sb = float((pk['sb'] or ('21 SB' if pk['cap'] == 'Ø100' else '4 SB')).split()[0])
        rows.append((f"P{j + 1}", pk['cap'], sb, 2 * sb, f"=ROUND(0.5*SQRT(D{5 + j}),2)", 'zemin asma tavanı → bodrum tavanı → ön cephe rögarı'))
    grid(ws, 4, ('Kolon', 'Çap', 'Kat SB', 'Toplam SB (2 kat)', 'Qww (l/s)', 'Bağlantı'), rows, (8, 8, 10, 16, 12, 56))
    # ---------------- havalandırma
    ws = wb.create_sheet('HAVALANDIRMA')
    head(ws, 'OTOPARK EGZOZ / TAZE HAVA / DUMAN TAHLİYESİ VE SIĞINAK HAVALANDIRMASI')
    o = R['otopark']; s = R['siginak']
    rows = [('Otopark net alan', o['A'], 'm²', 'tesisat alanı, sığınak, lobi ve çekirdekler hariç'),
            ('Otopark yüksekliği', o['h'], 'm', 'kabul'), ('Otopark hacmi', '=B5*B6', 'm³', ''),
            ('Normal havalandırma (4 hava değişimi)', '=B7*4', 'm³/h', 'kullanıcı talimatı'),
            ('Duman tahliyesi (10 hava değişimi)', '=B7*10', 'm³/h', 'alan > 100 m²'),
            ('Taze hava normal (%85)', '=ROUND(B8*0.85,0)', 'm³/h', 'otopark negatif basınçta'),
            ('Taze hava yangın (%85)', '=ROUND(B9*0.85,0)', 'm³/h', ''),
            ('Egzoz fanı (2 ad., çift hızlı, F300)', f"{o['Qd'] / 2:.0f} / {o['Qn'] / 2:.0f}", 'm³/h', 'her biri yüksek / düşük hız'),
            ('Sığınak alanı (ölçülen)', s['A'], 'm²', 'mimari iç ölçü 13,39 x 8,52'),
            ('Sığınak kişi sayısı', s['kisi'], 'kişi', '1 m²/kişi'),
            ('Sığınak hacmi (h=2,60)', '=B13*2.6', 'm³', ''),
            ('Normal havalandırma (4 hava değişimi)', '=ROUND(B15*4,0)', 'm³/h', ''),
            ('Koruma havası (3 m³/h·kişi, 51-150 kişi)', '=B14*3', 'm³/h', 'TS 3419 / kum filtreli'),
            ('Egzoz (+50 Pa için %80)', '=ROUND(B16*0.8,0)', 'm³/h', ''),
            ('Duman egzozu (alan > 100 m² → 10 hava değişimi)', '=ROUND(B15*10,0)' if s['duman'] else 'GEREKMEZ', 'm³/h', 'F400 fan')]
    grid(ws, 4, ('Kalem', 'Değer', 'Birim', 'Açıklama'), rows, (48, 16, 10, 48))
    wb.save(out)
    print(out)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
