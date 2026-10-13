# -*- coding: utf-8 -*-
"""VOYAP ofisleri: ısı kaybı (TS 2164 / örnek şablon, mahal bazlı) + yerden ısıtma çizelgeleri — 'örnek yerden ısıtma'
Excel şablonu (formüller korunur, yazdırma alanları ayarlı). Kaynak: yi.py önbelleği (çizimle aynı devreler)."""
import sys, datetime, math
sys.path.insert(0, '/home/claude/voyap/tes'); sys.path.append('/home/claude/nb'); sys.path.append('/home/claude/nb/p2')
import openpyxl
from openpyxl.utils import get_column_letter as CL
import yi
import gheat as H
import gexcel as X
from veri import *

KATLAR = [(2, '1 NK', 'ara', '1. KAT'), (3, '2 NK', 'ust', '2. KAT')]


def zno(kat, k, i):
    return f"{kat - 1}{k}-{i + 1:02d}"


def build(out):
    wb = openpyxl.load_workbook(X.SRC)
    ws = wb['kapak']
    ws['A2'] = ('ROTAŞ MÜHENDİSLİK.PROJE.MÜŞAVİRLİK.\nMak. Müh. Bayram GENÇ\n'
                'Semerciler Mah. Sait Faik Sk. No:22/C Z01 Adapazarı/Sakarya\nTEL: 0264 272 00 49   rotasmuhendislik@gmail.com\n')
    ws['B10'] = 'BAYRAM GENÇ'; ws['D10'] = None; ws['F10'] = None
    ws['A11'] = 'OFİSLER YERDEN ISITMA (MERKEZİ KAZAN DAİRESİ, 45/35 °C) ISI KAYBI VE DEVRE HESAP RAPORU'
    ws['C31'] = 'VOYAP - İŞ MERKEZİ (BODRUM + ZEMİN + 2 KAT, 5 DÜKKÂN + 28 OFİS)'
    ws['G41'] = datetime.datetime(2026, 10, 13); ws['H41'] = 'İŞ YERİ (DÜKKÂN + OFİS)'
    ws = wb['TEKNİK RAPOR']
    ws['E11'] = ': Merkezi sistem: sığınak içi kazan dairesi, duvar tipi yoğuşmalı kaskad kazanlar (doğalgaz)'
    ws['E12'] = ': Doğalgaz'
    ws['E13'] = ': Zeminden ısıtma borusu PEX 16x2,0 / 45-35 °C (3 yollu karışım vanalı ofis devresi)'
    ws['E14'] = ': Bitişik nizam'
    ws['E16'] = ': Bodrum + zemin (5 dükkân, radyatör) + 2 normal kat (her katta 14 ofis, yerden ısıtma)'
    ws['E17'] = ': Zeminden ısıtma borusu PEX 16x2,0 (oksijen bariyerli)'
    ws['E18'] = ': 45 - 35 °C'
    ws['F29'] = ':  -3°C (çatı)'
    ws.print_area = 'A1:N65'
    if 'TEKNİK RAPOR (2)' in wb.sheetnames: del wb['TEKNİK RAPOR (2)']
    tpl_wb = openpyxl.load_workbook(X.SRC)
    tpl = tpl_wb['ZEMİN']
    del wb['ZEMİN']
    ws1 = wb['1 NK']
    ws2 = wb.copy_worksheet(ws1); ws2.title = '2 NK'
    wb.move_sheet('2 NK', offset=-(len(wb.sheetnames) - 1 - wb.sheetnames.index('1 NK') - 1))
    TOT = {}
    DATA = {}
    for (kat, sh, kind, kad), wsx, no in zip(KATLAR, (ws1, ws2), (3, 4)):
        rooms = []
        for k in OFIS_S1:
            Y = yi.ofis_k(kat, k, None)
            DATA[(kat, k)] = Y
            for i, r in enumerate(Y['rooms']):
                c, Rr = yi.oda_calc(Y, r)
                e, _ = X.room_entries(Rr, kind, zno(kat, k, i))
                e = (e[0], f"OFİS {k} {r['name'][:14]}", *e[2:])
                rooms.append(e)
        TOT[kat] = X.write_heat_sheet(wsx, tpl, f"{kad} OFİSLER", no, rooms, str(kat - 1))
    # PETEK LİST
    ws = wb['PETEK LİST']
    X.clear_rows(ws, 10, 400, 16)
    tplr = [tpl_wb['PETEK LİST'].cell(10, c) for c in range(1, 17)]
    tpls = [tpl_wb['PETEK LİST'].cell(15, c) for c in range(1, 17)]
    r = 10
    PROW = {}
    for kat, sh, kind, kad in KATLAR:
        r0 = r
        for k in OFIS_S1:
            Y = DATA[(kat, k)]
            for i, rm in enumerate(Y['rooms']):
                z = zno(kat, k, i)
                cs = [c for c in Y['circ'] if c['room']['comp'] == rm['comp']]
                L_ = sum(c['L'] for c in cs)
                ti = yi.oda_calc(Y, rm)[0]['ti']
                vals = {'A': z, 'B': f"OFİS {k} {rm['name'][:14]}", 'C': str(ti), 'D': "='%s'!Q%d" % (sh, TOT[kat][z]), 'E': 17.5,
                        'F': "='YERDEN ISITMA'!N%d" % r, 'G': round(L_, 2), 'H': max(1, len(cs)) if cs else 0,
                        'I': 'ZEMİNDEN ISITMA' if cs else 'ISITILMIYOR (tefriş/ıslak hacim dolu) - kollektör bağlantı boruları'}
                for cc in range(1, 17):
                    X.copy_style(tplr[cc - 1], ws.cell(r, cc))
                for col, v in vals.items():
                    ws[col + str(r)] = v
                ws.merge_cells('I%d:P%d' % (r, r))
                PROW[(kat, k, i)] = (r, rm, cs)
                r += 1
        for cc in range(1, 17):
            X.copy_style(tpls[cc - 1], ws.cell(r, cc))
        ws['E%d' % r] = 'TOPLAM ='; ws['F%d' % r] = '=SUM(F%d:F%d)' % (r0, r - 1); ws['G%d' % r] = 'watt'
        r += 1
    ws['C1'] = 'MAHAL VE YERDEN ISITMA BORU BOYU ÇİZELGESİ'; ws['A1'] = 'ROTAŞ'
    ws['N3'] = datetime.datetime(2026, 10, 13)
    ws.print_area = 'A1:P%d' % r
    # YERDEN ISITMA
    ws = wb['YERDEN ISITMA']
    X.clear_rows(ws, 10, 400, 14)
    tplr = [tpl_wb['YERDEN ISITMA'].cell(10, c) for c in range(1, 15)]
    ws['A1'] = 'ROTAŞ'; ws['H6'] = 'BAYRAM GENÇ'
    for (kat, k, i), (rr, rm, cs) in PROW.items():
        s = cs[0]['s'] if cs else 15
        A = round(sum(c['Ar'] for c in cs), 2) if cs else rm['A']
        vals = {'A': "='PETEK LİST'!A%d" % rr, 'B': "='PETEK LİST'!B%d" % rr, 'C': "='PETEK LİST'!C%d" % rr,
                'D': "='PETEK LİST'!D%d" % rr, 'E': A, 'F': '=D%d/E%d' % (rr, rr), 'G': '=76.263-0.539*F%d' % rr,
                'H': '=D%d/G%d' % (rr, rr), 'I': s, 'J': "='PETEK LİST'!G%d" % rr, 'K': '=J%d/80' % rr,
                'L': len(cs), 'M': 17.5, 'N': '=J%d*M%d' % (rr, rr)}
        for cc in range(1, 15):
            X.copy_style(tplr[cc - 1], ws.cell(rr, cc))
        for col, v in vals.items():
            ws[col + str(rr)] = v
    ws.print_area = 'A1:N%d' % (max(v[0] for v in PROW.values()) + 2)

    # MMO (devre bazında)
    def mmo(ws, kat, title):
        X.clear_rows(ws, 7, 35, 200)
        tpl_col = [ws.cell(rr, 6) for rr in range(1, 48)]
        ws['A2'] = title
        allc = []
        for k in OFIS_S1:
            Y = DATA[(kat, k)]
            cl = []
            for c in sorted(Y['circ'], key=lambda z: z['pidx']):
                i = [j for j, rm in enumerate(Y['rooms']) if rm['comp'] == c['room']['comp']][0]
                cl.append(dict(port=c['pidx'] + 1, zone=zno(kat, k, i), room=f"{k} {c['room']['name'][:10]}", ti=c.get('ti', 20),
                               Ar=c['Ar'], W=max(c.get('W', 0.0), 1.0), s=c['s'], L=c['L']))
            allc += [(k, d) for d in X.hydraulics(cl)]
        ws['D3'] = 'Isıtma Çevrim Sayısı : %d' % len(allc)
        ws['D5'] = 'Mak. Müh. Bayram GENÇ'
        col = 6
        cols = []
        for k, d in allc:
            C = CL(col)
            v = {7: 'P%d' % d['port'], 8: 'Z' + d['zone'], 9: d['room'], 10: d['ti'], 11: 20,
                 12: round(d['Ar'], 2), 13: round(d['W'], 1), 14: '=%s13/%s12' % (C, C), 15: 0.012, 16: 1, 17: '=%s12' % C,
                 18: '=%s13/%s17' % (C, C), 19: '=%s10+(%s14/8.92)^0.91' % (C, C), 20: 45, 21: '=%s20-%s10' % (C, C),
                 22: d['s'], 23: 40, 24: '=%s23-%s10' % (C, C), 25: '=2*(%s21-%s24)' % (C, C),
                 26: '=0.2*%s14' % C, 27: '=%s17*(%s14+%s26)' % (C, C, C), 28: '=%s27/(1.163*%s25)' % (C, C),
                 29: '=%s17/(%s22*0.01)' % (C, C), 30: round(max(0.0, d['L'] - d['Ar'] / (d['s'] / 100.0)), 2),
                 31: '=%s29+%s30' % (C, C), 32: round(d['dP']), 33: '=%s32*0.05' % C, 34: '=%s32+%s33' % (C, C)}
            for rr, val in v.items():
                X.copy_style(tpl_col[rr - 1], ws.cell(rr, col))
                ws.cell(rr, col).value = val
            cols.append(C)
            col += 1
        for C in cols:
            X.copy_style(tpl_col[34], ws[C + '35'])
            ws[C + '35'] = '=MAX(%s34:%s34)-%s34' % (cols[0], cols[-1], C)
        ws['E2'] = 'TOPLAM ISI KAYBI = %.0f W' % sum(d['W'] for _, d in allc)
        ws['E3'] = 'MAX. BASINÇ KAYBI = %.0f Pa' % max(d['dP'] * 1.05 for _, d in allc)
        ws['C30'] = 'BAĞLANTI BORULARI'; ws['D30'] = 'Kat Planından (çizimden ölçülen)'
        ws['C31'] = 'TOPLAM BORU UZUNLUĞU'; ws['D31'] = 'Lç + Lb (çizimden ölçülen toplam)'
        for rr in range(36, 47):
            for cc in range(8, 300):
                ws.cell(rr, cc).value = None
        ws['B41'] = 'Kazan dairesinden 3 yollu karışım vanası ile gidiş 45 °C, dönüş 35 °C (ortalama 40 °C) seçilmiştir.'
        ws['B43'] = 'Alta geçen ısı 0,2 q" olarak alınmıştır.'
        ws['B44'] = 'R (Pa/m): PEX 16x2,0 (iç çap 12 mm), Blasius bağıntısı ile hesaplanmıştır.'
        ws['B46'] = 'Isıtma devrelerinin basıncı kollektördeki debimetreli reglaj vanası ile ayarlanmalıdır.'
        ws.print_area = 'A1:%s47' % cols[-1]
        ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 0; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    mmo(wb['MMO YERDEN ISITMA'], 2, '1. KAT OFİSLER (A1 - E2)')
    mmo(wb['YERDEN ISITMA- DEVAMI'], 3, '2. KAT OFİSLER (A1 - E2)')
    wb['YERDEN ISITMA- DEVAMI']['M5'] = None
    X.logo_temizle(wb)
    wb.save(out)
    return out


if __name__ == '__main__':
    print(build(sys.argv[1] if len(sys.argv) > 1 else '/tmp/claude-0/vy/out/VOYAP_ISI_KAYBI_YERDEN_ISITMA.xlsx'))
