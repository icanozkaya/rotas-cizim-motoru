# -*- coding: utf-8 -*-
import pickle, math, sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
sys.path.insert(0, '/home/claude/voyap')
import hidro
VY = '/tmp/claude-0/vy/'
res, dol = hidro.ozet()
R = pickle.load(open(VY + 'routes.pkl', 'rb')); H = pickle.load(open(VY + 'heads.pkl', 'rb'))
b = Side(style='thin'); B = Border(left=b, right=b, top=b, bottom=b)
HF = PatternFill('solid', fgColor='DDDDDD'); BOLD = Font(bold=True)
wb = openpyxl.Workbook()

def head(ws, t1, t2):
    ws['A1'] = 'ROTAŞ MÜHENDİSLİK - VOYAP - YANGIN TESİSATI (SPRİNKLER + YANGIN DOLABI)'; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = t2; ws['A2'].font = Font(bold=True, size=10)

def grid(ws, r0, hdr, rows, widths, fmt=None):
    for j, h in enumerate(hdr):
        c = ws.cell(r0, j + 1, h); c.font = BOLD; c.fill = HF; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='center')
    r = r0 + 1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(r, j + 1, v); c.border = B
            if fmt and j in fmt: c.number_format = fmt[j]
        r += 1
    for j, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = w
    return r

# ---- 1. tasarım esasları, pompa, depo
ws = wb.active; ws.title = 'POMPA-DEPO'
head(ws, '', 'TASARIM ESASLARI, POMPA VE DEPO HESABI (TS EN 12845, BYKHY md. 94, TS EN 671-1)')
w = max(res, key=lambda r: r['Q'])
rows = [
 ('En elverişsiz sprinkler alanı', 'Qsp', round(w['Q'], 1), 'l/dk', f"{w['kat']} {w['sinif']} (HİDROLİK sayfası)"),
 ('Yangın dolabı eş zamanlı adet', 'n', 2, 'adet', 'en uzak iki dolap'),
 ('Dolap debisi (TS EN 671-1, Ø25)', 'qd', 100, 'l/dk', 'BYKHY md. 94'),
 ('Dolap toplam debi', 'Qd', '=D5*D6', 'l/dk', ''),
 ('Toplam debi', 'Q', '=D4+D7', 'l/dk', ''),
 ('Toplam debi', 'Q', '=D8*0.06', 'm³/h', ''),
 ('Çalışma süresi (OH)', 't', 60, 'dk', 'TS EN 12845'),
 ('Gerekli su hacmi', 'V', '=D8*D10/1000', 'm³', ''),
 ('Mevcut depo net hacmi', 'Vd', 80, 'm³', '4,00 x 8,50 x 2,35 m'),
 ('Kontrol', '', '=IF(D12>=D11,"UYGUN","YETERSİZ")', '', ''),
 ('Sprinkler için pompa basıncı (en elverişsiz)', 'Psp', round(max(r['Pp'] for r in res) * 10.2, 1), 'mSS', 'HİDROLİK sayfası'),
 ('Dolap lüle basıncı', 'pl', 4.0, 'bar', 'BYKHY md. 94'),
 ('2. kat dolap kotu (pompa ekseninden)', 'z', 11.3, 'm', 'kat yükseklikleri kabulü'),
 ('Dolap hattı kayıpları (2" 25 m + 2½" 30 m + vana)', 'Δp', round(dol['P'] - 4.0 - 11.3 / 10.2, 2), 'bar', 'Hazen-Williams C=120'),
 ('Dolap için pompa basıncı', 'Pd', '=(D15+D17)*10.2+D16', 'mSS', ''),
 ('Gerekli pompa basıncı', 'H', '=MAX(D14,D18)', 'mSS', ''),
 ('Seçilen pompa (elektrikli ana + dizel yedek)', 'Qp', 70, 'm³/h', 'H=60 mSS'),
 ('Seçilen pompa basma yüksekliği', 'Hp', 60, 'mSS', ''),
 ('Kontrol: eğride Q=D9 için basınç ≥ D19', '', '=IF(AND(D21>=D19,D20*1.4>=D9),"UYGUN (eğriden teyit)","POMPA BÜYÜT")', '', '%140 debide ≥%70 basma'),
 ('Hidrolik güç', 'Ph', '=1000*9.81*D20/3600*D21/1000', 'kW', ''),
 ('Mil gücü (η=0,70)', 'Pm', '=D23/0.7', 'kW', ''),
 ('%150 debide güç (≈1,25 Pm)', 'P150', '=D24*1.25', 'kW', 'elektrikli 22 kW / dizel 30 kW'),
 ('Jokey pompa', 'Qj', 3, 'm³/h', 'H=70 mSS, 1,5 kW'),
 ('Dizel yakıt (3 saat, 0,25 l/kWh x 30 kW)', 'Vy', '=0.25*30*3', 'L', 'seçim 100 L'),
]
r = grid(ws, 3, ['NO', 'KALEM', 'SEMBOL', 'DEĞER', 'BİRİM', 'AÇIKLAMA'], [(k + 1,) + x for k, x in enumerate(rows)], (5, 52, 8, 14, 8, 40), {3: '0.00'})
ws.cell(r + 1, 1, 'NOT: Kat yükseklikleri bodrum 3,50 / zemin 3,50 / normal katlar 3,00 m kabul edilmiştir; mimari kesitten teyit edilecektir.')

# ---- 2. hidrolik
ws = wb.create_sheet('HİDROLİK')
head(ws, '', 'SPRİNKLER HİDROLİK HESABI - Hazen-Williams C=120, K80, pmin=0,35 bar, bağlantı payı %20 (ağaç üzerinde dengeli çözüm)')
rows = []
for r_ in res:
    rows.append((r_['kat'], r_['sinif'], r_['yog'], r_['alan'], r_['N'], r_['ara'], round(r_['qmin'], 1), round(r_['pdes'], 2), round(r_['Q'], 0), round(r_['P'], 2), round(r_['Pp'], 2), round(r_['Lrem'], 0)))
rr = grid(ws, 3, ['KAT', 'SINIF', 'mm/dk', 'Alan m²', 'Başlık', 'Ara sev.', 'q min l/dk', 'p uç bar', 'Q l/dk', 'P kat girişi bar', 'P pompa bar', 'En uzak yol m'], rows, (14, 8, 8, 9, 8, 8, 10, 9, 9, 13, 12, 12))
rr += 1
for r_ in res:
    ws.cell(rr, 1, f"{r_['kat']} - kat girişinden en uzak başlığa boru yolu").font = BOLD; rr += 1
    hdr = ['Çap', 'DN', 'İç çap mm', 'L m', 'Q l/dk', 'Δp bar', 'v m/s']
    for j, h in enumerate(hdr):
        c = ws.cell(rr, j + 1, h); c.font = BOLD; c.fill = HF; c.border = B
    rr += 1
    r1 = rr
    for c_, dn, L_, q in r_['det']:
        vals = [c_, dn, hidro.ID[dn], round(L_, 2), round(q, 1),
                f'=6.05E5*E{rr}^1.85/(120^1.85*C{rr}^4.87)*D{rr}*1.2', f'=E{rr}/60000/(PI()*(C{rr}/1000)^2/4)']
        for j, v in enumerate(vals):
            c = ws.cell(rr, j + 1, v); c.border = B
            if j in (5, 6): c.number_format = '0.000'
        rr += 1
    ws.cell(rr, 5, 'Toplam'); ws.cell(rr, 6, f'=SUM(F{r1}:F{rr - 1})').number_format = '0.000'
    rr += 2

# ---- 3. boru metrajı
ws = wb.create_sheet('METRAJ')
head(ws, '', 'BORU VE MALZEME METRAJI (%10 fire hariç)')
caps = ['1"', '1 1/4"', '1 1/2"', '2"', '2 1/2"', '3"', '4"']
L = {i: {c: 0.0 for c in caps} for i in range(4)}
for i in range(4):
    for e in R[i]['ts']['edges']: L[i][e['cap']] += e['L'] / 100.0
rows = []
for c in caps:
    rows.append((c,) + tuple(round(L[i][c], 1) for i in range(4)) + (f'=SUM(B{len(rows) + 4}:E{len(rows) + 4})',))
rr = grid(ws, 3, ['Çap', 'BODRUM', 'ZEMİN', '1. KAT', '2. KAT', 'TOPLAM m'], rows, (12, 12, 12, 12, 12, 12), {5: '0.0'})
rr += 1
nU = sum(1 for i in (0, 1) for h in H[i] if h['t'] == 'U'); nP = sum(1 for i in (2, 3) for h in H[i]); nA = sum(1 for h in H[0] if h['t'] == 'ARA')
items = [('Sprinkler upright K80 68°C', nU, 'adet'), ('Sprinkler pendant K80 68°C', nP, 'adet'), ('Ara seviye sprinkler (lift)', nA, 'adet'),
         ('Yangın dolabı TS EN 671-1', 8, 'adet'), ('Islak alarm vanası DN100', 2, 'adet'), ('Kat kontrol istasyonu DN100', 3, 'takım'),
         ('İtfaiye bağlantısı 2x2½"', 1, 'adet'), ('Yangın pompa grubu (E+D+J)', 1, 'takım'), ('Yangın suyu deposu 80 m³', 1, 'adet'),
         ('Yedek sprinkler dolabı (6 adet başlık + anahtar)', 1, 'adet')]
grid(ws, rr, ['MALZEME', 'MİKTAR', 'BİRİM'], items, (40, 12, 10))

out = VY + 'out/VOYAP_YANGIN_HESAP.xlsx'
import os; os.makedirs(VY + 'out', exist_ok=True)
wb.save(out); print(out)
