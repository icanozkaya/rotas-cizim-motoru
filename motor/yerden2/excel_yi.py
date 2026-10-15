"""yerden2 — hesap Excel'i: sablon/ornek_yi.xlsx kopyası; şablon formülleri korunur, yalnız girdiler yazılır.

Sayfalar:
  kapak / TEKNİK RAPOR   : proje bilgisi, rejim (45/35 °C), mahal sıcaklıkları
  PETEK LİST             : mahal satırları (zon, ad, sıcaklık, ısıl güç, boru boyu, devre sayısı); F = 'YERDEN ISITMA'!N
  YERDEN ISITMA          : modülasyon (şablon formülleri F, G, H, K, N aynen); E = mahal m² (mimari yazısı)
  MMO YERDEN ISITMA / YERDEN ISITMA- DEVAMI / ... (2) : devre bazında MMO özeti, kat başına bir sayfa
Isı kaybı sayfaları (ZEMİN, 1 NK) ve TEKNİK RAPOR (2) şablondaki BAŞKA projenin verisidir -> kaldırılır
(bu projede ısı kaybı hesabı yok; "D" sütunu çizimdeki 35 W/m × L değeridir — BENIOKU'da belirtilir).
"""
import copy
import datetime
import math
import os
import shutil

import openpyxl
from openpyxl.utils import get_column_letter as CL

BIRIM = 35.0      # W/m — plan etiketi, kolon şeması ve Excel AYNI esas: devre ısıl gücü = 35 W/m × L (R6)
TS_SINIR = {20: 29.0, 15: 29.0, 24: 33.0}   # MMO döşeme yüzey sıcaklığı sınırı (oturma 29, banyo 33 °C)
KOK = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SABLON = KOK + '/sablon/ornek_yi.xlsx'
CIKIS = KOK + '/proje/denemeler/yerden2/NB_YERDEN_ISITMA.xlsx'


def stil(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.border = copy.copy(src.border)
        dst.fill = copy.copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy.copy(src.alignment)


def temizle(ws, r0, r1, c1):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= r0 and rng.max_row <= r1:
            ws.unmerge_cells(str(rng))
    for r in range(r0, r1 + 1):
        for c in range(1, c1 + 1):
            ws.cell(r, c).value = None


def dp_devre(W, L_m, dT=10.0):
    """PEX 16x2,0 (di 12 mm), 40 °C su; Blasius / laminer -> (debi kg/h, R Pa/m, ΔPs Pa)"""
    m = W / (1.163 * dT)
    rho, nu, di = 992.0, 0.658e-6, 0.012
    Q = m / rho / 3600.0
    v = Q / (math.pi * di * di / 4)
    Re = v * di / nu
    lam = 64.0 / Re if Re < 2300 else 0.3164 * Re ** -0.25
    R = lam / di * rho * v * v / 2
    return m, R, R * L_m


YUZEY = []


def yaz(model, cikis=CIKIS):
    shutil.copyfile(SABLON, cikis)
    YUZEY.clear()
    wb = openpyxl.load_workbook(cikis)
    tpl = openpyxl.load_workbook(SABLON)
    P = model['proje']
    katlar = model['katlar']
    mahaller = []      # (kat, daire, mahal)
    for kat in katlar:
        for dn in sorted(kat['daireler'], key=int):
            for m in kat['daireler'][dn]['mahaller']:
                mahaller.append((kat, kat['daireler'][dn]['no'], m))
    # ---------------- kapak
    ws = wb['kapak']
    ws['A2'] = ('ROTAŞ MÜHENDİSLİK.PROJE.MÜŞAVİRLİK.\n'
                'Semerciler Mah. Sait Faik Sk. No:22/C Z01 Adapazarı/Sakarya\nrotasmuhendislik@gmail.com\n')
    ws['B10'] = 'CAN ÖZKAYA'
    ws['D10'] = None
    ws['F10'] = None
    ws['A11'] = 'YERDEN ISITMA TESİSATI PROJE RAPORU'
    ws['C31'] = P['ad']
    for c in ('A36', 'B36', 'D36', 'G36', 'H36'):
        ws[c] = '-'
    ws['I36'] = '4739/4'
    ws['A41'] = 'BETONARME'
    ws['C41'] = 9
    ws['D41'] = 'ZEMİN + 2'
    ws['G41'] = datetime.datetime(2026, 10, 15)
    ws['H41'] = 'KONUT'
    # ---------------- teknik rapor
    ws = wb['TEKNİK RAPOR']
    ws['A1'] = 'YERDEN ISITMA TESİSATI PROJESİ TEKNİK RAPORU'
    ws['F8'] = None
    ws['E10'] = ': Konut (Z+2, 9 daire)'
    ws['E11'] = ': Duvar tipi yoğuşmalı kombi (her dairede 1 adet, 24 kW)'
    ws['E13'] = ': Zeminden ısıtma borusu PEX 16x2,0 / 45-35 °C, boru aralığı 15 cm'
    ws['E16'] = ': Zemin + 2 normal kat'
    ws['E17'] = ': Zeminden ısıtma borusu PEX 16x2,0 (oksijen bariyerli)'
    ws['E18'] = ': 45 - 35 °C'
    ws['B21'] = 'Salon'; ws['F21'] = ': 20°C'
    ws['B22'] = 'Oda / Yatak Odası'; ws['F22'] = ': 20°C'
    ws['B23'] = 'Mutfak'; ws['F23'] = ': 20°C'
    ws['B24'] = 'Banyo'; ws['F24'] = ': 24°C'
    ws['B25'] = 'Antre - Koridor'; ws['F25'] = ': 15°C (ayrı devre)'
    for c in ('F26', 'F27', 'F28'):
        ws[c] = None
    ws['B31'] = 'Isıtılmayan'; ws['F31'] = ': balkon, yüklük, wc-lavabo / duş / wc-duş (≤ 2,7 m²), merdiven, asansör'
    # 1.3: şablondaki U değerleri başka projeye ait (ısı kaybı sayfaları yok) -> ısıl güç esası yazılır (R6)
    for rr in range(33, 70):
        for cc in range(1, 15):
            ws.cell(rr, cc).value = None
    ws['A33'] = '1.3.'
    ws['B33'] = 'Devre Isıl Gücü Esası:'
    ws['B35'] = ('Plan etiketlerindeki, kolon şemasındaki ve bu hesaptaki devre ısıl gücü aynı esastan alınmıştır:')
    ws['B36'] = '   Q (W) = 35 W/m × hat uzunluğu L (m); L = plandaki serpantinin yaylar dahil gerçek boyu.'
    ws['B37'] = 'Bu projede ayrı ısı kaybı hesabı yapılmamıştır; "ısı ihtiyacı" sütunları boru boyundan hesaplanmıştır.'
    ws['B38'] = ('Boru PEX 16 x 2,0, boru aralığı 15 cm, duvardan kenar mesafesi 7,5 cm (dar ıslak hacimlerde 5-6 cm); '
                 'devre boyu en çok 80 m.')
    ws['B39'] = 'Döşeme yüzey sıcaklığı kontrolü MMO sayfalarındadır (sınır: oturma 29 °C, banyo 33 °C).'
    # K-1 (tur 3): esas değişmedi; şablon formülünün bu esasla verdiği anlamsız değerlerin nedeni not edilir
    ws['A41'] = '1.4.'
    ws['B41'] = 'Not — "YERDEN ISITMA" (modülasyon) sayfasındaki negatif değerler:'
    ws['B43'] = ('Sayfadaki G sütunu (MODÜL) şablon formülüdür: G = 76,263 − 0,539 × F, F = q = Q / A (W/m²). Şablon '
                 '17,5 W/m esası için kurulmuştur;')
    ws['B44'] = ('q > 141,5 W/m² olduğunda G negatif çıkar. Bu projede Q = 35 W/m × L ve boru yoğunluğu ≈ 6 m/m² '
                 'olduğundan q ≈ 200-215 W/m² olur;')
    ws['B45'] = ('bu yüzden G (MODÜL) ve H = D / G (m²) sütunlarında negatif değerler görünür. Bunlar bir ölçü değildir, '
                 'yalnız formülün geçerlilik aralığı dışında olunduğunu gösterir.')
    ws['B46'] = ('Formüller bozulmadan bırakılmıştır. Aynı esasla MMO sayfasında yüzey sıcaklığı ≈ 37,8 °C ve devre basınç '
                 'kaybı Δpmax üstünde çıkar (sayfadaki uyarı satırları).')
    ws['B47'] = ('Devre gücü esası (35 W/m × L, plan etiketleriyle aynı) proje sahibinin kararına bırakılmıştır; '
                 '17,5 W/m ya da ısı kaybı esasına geçilirse bu değerler olağan aralığa döner.')
    for nm in ('ZEMİN', '1 NK', 'TEKNİK RAPOR (2)'):
        if nm in wb.sheetnames:
            del wb[nm]
    # ---------------- PETEK LİST
    ws = wb['PETEK LİST']
    temizle(ws, 10, 200, 16)
    t_s = [tpl['PETEK LİST'].cell(10, c) for c in range(1, 17)]
    t_t = [tpl['PETEK LİST'].cell(15, c) for c in range(1, 17)]
    ws['A1'] = 'ROTAŞ'
    ws['C1'] = 'MAHAL VE YERDEN ISITMA BORU BOYU ÇİZELGESİ'
    ws['N3'] = datetime.datetime(2026, 10, 15)
    satir = {}
    r = 10
    for kat in katlar:
        r0 = r
        for k2, dno, m in mahaller:
            if k2 is not kat:
                continue
            L = sum(c['L_m'] for c in m['devreler'])
            W = BIRIM * L
            v = {'A': m['zon'], 'B': '%s %s' % (dno, m['ad'].upper()), 'C': str(m['sic']), 'D': round(W, 1),
                 'E': BIRIM, 'F': "='YERDEN ISITMA'!N%d" % r, 'G': round(L, 2), 'H': len(m['devreler']),
                 'I': 'ZEMİNDEN ISITMA'}
            for c in range(1, 17):
                stil(t_s[c - 1], ws.cell(r, c))
            for col, val in v.items():
                ws[col + str(r)] = val
            ws.merge_cells('I%d:P%d' % (r, r))
            satir[(kat['kat'], m['zon'])] = r
            r += 1
        for c in range(1, 17):
            stil(t_t[c - 1], ws.cell(r, c))
        ws['B%d' % r] = kat['ad']
        ws['E%d' % r] = 'TOPLAM ='
        ws['F%d' % r] = '=SUM(F%d:F%d)' % (r0, r - 1)
        ws['G%d' % r] = 'watt'
        r += 1
    ws.print_area = 'A1:P%d' % r
    ws['D6'] = 'ISIL GÜÇ\n35 W/m × L'
    ws['E6'] = 'BİRİM \nVERİMİ'
    ws['D9'] = 'W'
    ws['E9'] = 'W/m'
    ws['F9'] = 'W'
    # ---------------- YERDEN ISITMA (modülasyon) — şablon formülleri
    ws = wb['YERDEN ISITMA']
    temizle(ws, 10, 200, 14)
    t_s = [tpl['YERDEN ISITMA'].cell(10, c) for c in range(1, 15)]
    ws['A1'] = 'ROTAŞ'
    ws['H6'] = 'CAN ÖZKAYA'
    for k2, dno, m in mahaller:
        rr = satir[(k2['kat'], m['zon'])]
        v = {'A': "='PETEK LİST'!A%d" % rr, 'B': "='PETEK LİST'!B%d" % rr, 'C': "='PETEK LİST'!C%d" % rr,
             'D': "='PETEK LİST'!D%d" % rr, 'E': m.get('alan_mimari_m2') or m['alan_m2'], 'F': '=D%d/E%d' % (rr, rr),
             'G': '=76.263-0.539*F%d' % rr, 'H': '=D%d/G%d' % (rr, rr), 'I': 15, 'J': "='PETEK LİST'!G%d" % rr,
             'K': '=J%d/80' % rr, 'L': '=CEILING(K%d,1)' % rr, 'M': BIRIM, 'N': '=J%d*M%d' % (rr, rr)}
        for c in range(1, 15):
            stil(t_s[c - 1], ws.cell(rr, c))
        for col, val in v.items():
            ws[col + str(rr)] = val
    ws.print_area = 'A1:N%d' % (max(satir.values()) + 2)
    ws['D8'] = 'Mahal ısı ihtiyacı Q (boru boyundan: 35 W/m × L)'
    ws['M9'] = 'W/m'
    # ---------------- MMO: kat başına bir sayfa
    adlar = ['MMO YERDEN ISITMA', 'YERDEN ISITMA- DEVAMI']
    while len(adlar) < len(katlar):
        yeni = wb.copy_worksheet(wb['YERDEN ISITMA- DEVAMI'])
        yeni.title = 'YERDEN ISITMA- DEVAMI (%d)' % len(adlar)
        adlar.append(yeni.title)
    tpl_col = [tpl['MMO YERDEN ISITMA'].cell(rr, 6) for rr in range(1, 48)]
    for kat, ad in zip(katlar, adlar):
        ws = wb[ad]
        temizle(ws, 7, 35, 80)
        for rr in range(2, 6):
            for cc in range(4, 80):
                if ws.cell(rr, cc).value and cc != 3:
                    ws.cell(rr, cc).value = None
        ws['A2'] = kat['ad']
        ws['C2'] = 'DÖŞEMEDEN ISITMA HESAP ÖZETİ'
        ws['D2'] = 'Ts.g  = 45 °C'
        allc = []
        for k2, dno, m in mahaller:
            if k2 is not kat:
                continue
            Ltop = sum(c['L_m'] for c in m['devreler']) or 1.0
            for c in m['devreler']:
                allc.append((dno, m, c, m['isitma_alani_m2'] * c['L_m'] / Ltop))
        ws['D3'] = 'Isıtma Çevrim Sayısı : %d' % len(allc)
        ws['D5'] = 'Mak. Müh. CAN ÖZKAYA'
        col = 6
        cols = []
        maxdp = 0.0
        for dno, m, c, A in allc:
            C = CL(col)
            Wc = BIRIM * c['L_m']
            debi, R, dps = dp_devre(Wc, c['L_m'])
            maxdp = max(maxdp, dps * 1.05)
            v = {7: c['no'], 8: m['zon'], 9: '%s %s' % (dno, m['ad']), 10: m['sic'],
                 11: 6 if kat['kat'] == 'Z' else 20, 12: round(A, 2), 13: round(Wc, 1),
                 14: '=%s13/%s12' % (C, C), 15: 0.012, 16: 1, 17: '=%s12' % C, 18: '=%s13/%s17' % (C, C),
                 19: '=%s10+(%s14/8.92)^0.91' % (C, C), 20: 45, 21: '=%s20-%s10' % (C, C), 22: 15, 23: 40,
                 24: '=%s23-%s10' % (C, C), 25: '=2*(%s21-%s24)' % (C, C), 26: '=0.2*%s14' % C,
                 27: '=%s17*(%s14+%s26)' % (C, C, C), 28: '=%s27/(1.163*%s25)' % (C, C),
                 29: '=%s17/(%s22*0.01)' % (C, C), 30: 0, 31: '=%s29+%s30' % (C, C), 32: round(dps),
                 33: '=%s32*0.05' % C, 34: '=%s32+%s33' % (C, C)}
            for rr, val in v.items():
                stil(tpl_col[rr - 1], ws.cell(rr, col))
                ws.cell(rr, col).value = val
            cols.append(C)
            col += 1
        for C in cols:
            stil(tpl_col[34], ws[C + '35'])
            ws[C + '35'] = '=MAX(%s34:%s34)-%s34' % (cols[0], cols[-1], C)
        ws['E2'] = 'TOPLAM ISI = %.0f W' % sum(BIRIM * c['L_m'] for _, _, c, _ in allc)
        ws['E3'] = 'MAX. BASINÇ KAYBI = %.0f Pa' % maxdp
        ws['C30'] = 'BAĞLANTI BORULARI'
        ws['D30'] = 'Kollektör ayakları planda çizilmedi (0)'
        ws['C32'] = 'DÜZ BORU BASINÇ KAYBI (5)'
        ws['D32'] = 'R.L (PEX 16x2,0, çizimden ölçülen L)'
        for rr in range(36, 47):
            for cc in range(8, 80):
                ws.cell(rr, cc).value = None
        ws['B41'] = 'Kombi gidiş suyu sıcaklığı 45 °C, dönüş 35 °C seçilmiştir.'
        ws['B42'] = 'Ortalama su sıcaklığı 40 °C alınmıştır.'
        ws['B43'] = ('Alta geçen ısı 0,2 q" olarak alınmıştır. qh = 35 W/m × devre boyu (plan etiketi ve kolon şeması ile '
                     'aynı esas); ayrı ısı kaybı hesabı yapılmamıştır.')
        ws['C13'] = 'DEVRE ISIL GÜCÜ'
        ws['D13'] = '35 W/m × L (planla aynı)'
        # yüzey sıcaklığı kontrolü (Ty = Ti + (q"/8,92)^0,91; sınır oturma 29 °C, banyo 33 °C)
        from openpyxl.styles import PatternFill, Font
        kirmizi = PatternFill('solid', fgColor='FFC7CE')
        asim = []
        for C, (dno, m, c, A) in zip(cols, allc):
            q = BIRIM * c['L_m'] / A if A > 0 else 0.0
            ty = m['sic'] + (q / 8.92) ** 0.91
            sinir = TS_SINIR.get(m['sic'], 29.0)
            YUZEY.append(dict(kat=kat['kat'], zon=m['zon'], devre=c['no'], ad=m['ad'], sic=m['sic'], A=round(A, 2),
                              L=c['L_m'], q=round(q, 1), Ty=round(ty, 2), sinir=sinir))
            if ty > sinir:
                ws[C + '19'].fill = kirmizi
                asim.append('%s-%d (%.1f °C)' % (m['zon'], c['no'], ty))
        if asim:
            ws['B38'] = ('UYARI: 35 W/m × L esasında döşeme yüzey sıcaklığı sınırı aşılıyor (%d devre): %s' %
                         (len(asim), ', '.join(asim[:8]) + (' ...' if len(asim) > 8 else '')))
            ws['B39'] = ('Sınır içinde kalmak için q" ≤ 100 W/m² (oturma) gerekir; 35 W/m ile bu, boru aralığı ≈ 35 cm '
                         'demektir. Boru aralığı 15 cm (kullanıcı seçimi) değiştirilmemiştir — ısı kaybı ile qh girilmelidir.')
            ws['B38'].font = Font(bold=True, color='C00000')
            ws['B39'].font = Font(color='C00000')
        ws['B44'] = 'R (Pa/m): PEX 16x2,0 (iç çap 12 mm), 40 °C su, Blasius / laminer bağıntısı; L = çizimdeki devre boyu.'
        ws['B46'] = 'Isıtma devrelerinin basıncı kollektördeki debimetreli reglaj vanası ile ayarlanmalıdır.'
        ws['A12'] = 6
        ws.print_area = 'A1:%s47' % (cols[-1] if cols else 'F')
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    # yabancı logo / resim temizliği (dcm makina)
    for w in wb.worksheets:
        w._images = []
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'dcm' in c.value.lower():
                    c.value = 'ROTAŞ'
    # R9: şablondan kalan dış bağlantı (…/dcm/PROJE/MEKANİK HESAPLAR KOMBİ.xls) ve tüm 'dcm' izleri silinir
    wb._external_links = []
    for w in wb.worksheets:
        for hf in (w.oddHeader, w.oddFooter, w.evenHeader, w.evenFooter, w.firstHeader, w.firstFooter):
            for part in (hf.left, hf.center, hf.right):
                if part.text and 'dcm' in part.text.lower():
                    part.text = None
    for dn in list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else []:
        d = wb.defined_names[dn]
        if d.attr_text and ('[' in d.attr_text or 'dcm' in d.attr_text.lower()):
            del wb.defined_names[dn]
    wb.properties.creator = 'ROTAŞ MÜHENDİSLİK - CAN ÖZKAYA'
    wb.properties.lastModifiedBy = 'ROTAŞ MÜHENDİSLİK'
    wb.properties.title = 'YERDEN ISITMA HESAPLARI'
    # R9: kapak ve TEKNİK RAPOR yazdırma alanı (A4 dikey, sayfaya sığdır)
    for ad, alan in (('kapak', 'A1:I42'), ('TEKNİK RAPOR', 'A1:H48')):
        w = wb[ad]
        w.print_area = alan
        w.page_setup.orientation = 'portrait'
        w.page_setup.paperSize = 9
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 1
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(cikis)
    _dcm_temizle(cikis)
    return cikis


def _dcm_temizle(yol):
    """kaydedilmiş paketi tara: externalLink parçası / 'dcm' izi kaldıysa zip'ten çıkar (R9)"""
    import zipfile
    import re
    gecici = yol + '.tmp'
    with zipfile.ZipFile(yol) as zi, zipfile.ZipFile(gecici, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            if it.filename.startswith('xl/externalLinks/'):
                continue
            veri = zi.read(it.filename)
            if it.filename in ('[Content_Types].xml', 'xl/_rels/workbook.xml.rels', 'xl/workbook.xml'):
                t = veri.decode('utf-8')
                t = re.sub(r'<Override[^>]*externalLink[^>]*/>', '', t)
                t = re.sub(r'<Relationship[^>]*externalLink[^>]*/>', '', t)
                t = re.sub(r'<externalReferences>.*?</externalReferences>', '', t)
                veri = t.encode('utf-8')
            zo.writestr(it, veri)
    os.replace(gecici, yol)
