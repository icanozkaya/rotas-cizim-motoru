"""Doğalgaz hesap Excel'i (openpyxl, formüllü).

Sayfalar: PARAMETRE, BORU, ESZAMANLILIK, CIHAZLAR, CAP_HESABI, KONTROL, HAVALANDIRMA.
CAP_HESABI'nda DN değeri motorun seçtiği değerdir (elle değiştirilebilir); hız, kayıplar ve kontroller formüldür.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

import hesap

INCE = Side(style='thin', color='808080')
KEN = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
BAS = PatternFill('solid', fgColor='FFF2CC')
KIR = PatternFill('solid', fgColor='F8CBAD')
YES = PatternFill('solid', fgColor='C6EFCE')
B = Font(bold=True)
AKSA = 'AKSAGAZ şartnamesiyle teyit edilecek'


def _baslik(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c)
        x.font, x.fill, x.border = B, BAS, KEN
        x.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')


def _ken(ws, r0, r1, c1):
    for r in range(r0, r1 + 1):
        for c in range(1, c1 + 1):
            ws.cell(r, c).border = KEN


def yaz(ag, P, yol):
    wb = Workbook()
    # ------------------------------------------------------------------ PARAMETRE
    ws = wb.active
    ws.title = 'PARAMETRE'
    ws['A1'] = ('DENEME - ' if P.get('deneme') else '') + P['proje']
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = f'Dağıtım şirketi: {P["dagitim_sirketi"]} - değer ve sınırlar {AKSA}.'
    ws['A2'].font = Font(bold=True, color='C00000')
    par = [
        ('Sistem basıncı', P['sistem_basinci_mbar'], 'mbar', 'P_sis'),
        ('Alt ısıl değer Hu', P['gaz']['Hu_kcal_m3'], 'kcal/m³', 'Hu'),
        ('Cihaz verimi η', P['gaz']['verim'], '-', 'eta'),
        ('Bağıl yoğunluk d', P['gaz']['bagil_yogunluk'], '-', 'd_gaz'),
        ('Servis kutusu → sayaç vanası sınırı', P['sinirlar']['servis_sayac'], 'mbar', 'lim_on'),
        ('Sayaç → cihaz sınırı', P['sinirlar']['sayac_cihaz'], 'mbar', 'lim_so'),
        ('Toplam (sayaç hariç) sınırı', P['sinirlar']['toplam'], 'mbar', 'lim_top'),
        ('Hız sınırı', P['sinirlar']['hiz'], 'm/s', 'vmax'),
        ('Bağımsız bölüm min. debisi', P['eszamanlilik']['bb_min_debi'], 'm³/h', 'qbb_min'),
        ('kcal/h → kW', 860, 'kcal/kWh', 'kcal_kw'),
    ]
    _baslik(ws, 4, ['Parametre', 'Değer', 'Birim', 'Ad'])
    from openpyxl.workbook.defined_name import DefinedName
    for i, (a, v, u, n) in enumerate(par, 5):
        ws.cell(i, 1, a), ws.cell(i, 2, v), ws.cell(i, 3, u), ws.cell(i, 4, n)
        wb.defined_names[n] = DefinedName(n, attr_text=f"PARAMETRE!$B${i}")
    _ken(ws, 4, 4 + len(par), 4)
    r = 6 + len(par)
    notlar = [
        'Formüller (kutuphane/arastirma/07_ISITMA_DOGALGAZ.md §6.3):',
        'Cihaz debisi Q = kW·860/(Hu·η) [m³/h]; ocak için katalog debisi (1,6 m³/h).',
        'Sürtünme ΔP_R = 23200·d·L·Q^1,82·D^-4,82 [mbar] (Renouard, alçak basınç, D iç çap mm).',
        'Yerel kayıp ΔP_Z = 3,97·10⁻³·Σξ·V² [mbar]; ξ: 90° dirsek 0,4 / T ayrılma 1,3 / küresel vana 0,5.',
        'Yükseklik ΔP_H = -0,049·h [mbar] (yükselişte kazanç, inişte kayıp).',
        'Hız V = Q/(3600·π·D²/4) [m/s].',
        'Eşzamanlılık katsayısı tablosu ' + AKSA + ' (ESZAMANLILIK sayfası).',
        'Kombi (C tipi) 150 cm² dış hava menfezi; ocak (A tipi) mahal ≥12 m³ ve ≥1 m³/kW (HAVALANDIRMA sayfası).',
    ]
    for s in notlar:
        ws.cell(r, 1, s)
        r += 1
    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 12
    # ------------------------------------------------------------------ BORU
    wb_b = wb.create_sheet('BORU')
    _baslik(wb_b, 1, ['DN', 'Dış çap (mm)', 'Et (mm)', 'İç çap (mm)', 'İnç'])
    for i, b in enumerate(hesap.BORU, 2):
        for j, v in enumerate(b, 1):
            wb_b.cell(i, j, v)
    _ken(wb_b, 1, len(hesap.BORU) + 1, 5)
    wb_b['G1'] = 'TS EN 10255 orta seri çelik boru'
    # ------------------------------------------------------------------ ESZAMANLILIK
    we = wb.create_sheet('ESZAMANLILIK')
    _baslik(we, 1, ['n (BB)', 'f'])
    for n in range(1, 151):
        we.cell(n + 1, 1, n)
        we.cell(n + 1, 2, hesap.esz_f(n, P['eszamanlilik']))
    we['D1'] = 'Eşzamanlılık katsayısı - ' + AKSA + '. Tablo boşken f = 1,0 (emniyetli taraf); >100 BB f = 0,397.'
    we['D1'].font = Font(bold=True, color='C00000')
    # ------------------------------------------------------------------ CIHAZLAR
    wc = wb.create_sheet('CIHAZLAR')
    _baslik(wc, 1, ['Bölüm', 'Kat', 'Cihaz', 'Tip', 'Kapasite (kW)', 'Katalog debisi (m³/h)', 'Debi Q (m³/h)',
                    'Bölüm toplamı (m³/h)', 'Bölüm hesap debisi (m³/h)'])
    r = 2
    cih_ref = {}      # (daire, cihaz) -> hücre
    bb_ref = {}       # daire -> bölüm hesap debisi hücresi
    for d in ag.daireler:
        r0 = r
        for c in d['cihazlar']:
            wc.cell(r, 1, d['no'])
            wc.cell(r, 2, P['katlar'][d['kat']])
            wc.cell(r, 3, c['ad'])
            wc.cell(r, 4, c['tip'])
            if c.get('debi'):
                wc.cell(r, 6, c['debi'])
                wc.cell(r, 5, f'=F{r}*Hu/kcal_kw')
                wc.cell(r, 7, f'=F{r}')
            else:
                wc.cell(r, 5, c['kW'])
                wc.cell(r, 7, f'=E{r}*kcal_kw/(Hu*eta)')
            cih_ref[(d['no'], c['ad'])] = f'CIHAZLAR!$G${r}'
            r += 1
        wc.cell(r0, 8, f'=SUM(G{r0}:G{r - 1})')
        wc.cell(r0, 9, f'=MAX(H{r0},qbb_min)')
        bb_ref[d['no']] = f'CIHAZLAR!$I${r0}'
    _ken(wc, 1, r - 1, 9)
    for col, w in zip('ABCDEFGHI', (8, 12, 10, 6, 14, 18, 14, 18, 20)):
        wc.column_dimensions[col].width = w
    # ------------------------------------------------------------------ CAP_HESABI
    wh = wb.create_sheet('CAP_HESABI')
    cols = ['Hat', 'Baş.', 'Son', 'Açıklama', 'Tür', 'Cihazlar', 'n (BB)', 'Q top (m³/h)', 'f', 'Q hesap (m³/h)',
            'L (m)', 'h (m)', 'DN', 'Di (mm)', 'V (m/s)', 'ΔP/L (mbar/m)', 'ΔP_R (mbar)', 'Σξ', 'ΔP_Z (mbar)',
            'ΔP_H (mbar)', 'ΔP_Σ (mbar)', 'Kümülatif ΔP (mbar)', 'Hız kontrolü']
    wh['A1'] = 'DOĞALGAZ BORU ÇAPI VE BASINÇ KAYBI HESABI (' + str(P['sistem_basinci_mbar']) + ' mbar) - ' + AKSA
    wh['A1'].font = Font(bold=True, size=12)
    _baslik(wh, 3, cols)
    satir = {}
    by_b = {h.b: h for h in ag.hatlar}
    for i, h in enumerate(ag.hatlar, 4):
        satir[h.no] = i
    dmap = {d['no']: d for d in ag.daireler}
    for h in ag.hatlar:
        i = satir[h.no]
        wh.cell(i, 1, h.no)
        wh.cell(i, 2, h.a)
        wh.cell(i, 3, h.b)
        wh.cell(i, 4, h.aciklama)
        wh.cell(i, 5, h.tur)
        wh.cell(i, 6, ', '.join(h.cihazlar))
        wh.cell(i, 7, h.n_bb)
        if h.tur == 'cihaz':
            wh.cell(i, 8, '=' + cih_ref[(h.bolum, h.cihaz['ad'])])
        elif h.tur in ('daire', 'sayac_baglanti'):
            wh.cell(i, 8, '=' + bb_ref[h.bolum])
        else:
            if h.tur == 'kollektor':
                k = ag.kollektor.index(h)
                akis = ag.daireler[k + 1:]
            else:
                akis = ag.daireler
            wh.cell(i, 8, '=' + '+'.join(bb_ref[d['no']] for d in akis))
        if h.tur == 'cihaz':
            wh.cell(i, 9, 1)
        else:
            wh.cell(i, 9, f'=VLOOKUP(G{i},ESZAMANLILIK!$A$2:$B$151,2,TRUE)')
        wh.cell(i, 10, f'=H{i}*I{i}')
        wh.cell(i, 11, h.L)
        wh.cell(i, 12, h.h)
        wh.cell(i, 13, h.dn)
        wh.cell(i, 14, f'=VLOOKUP(M{i},BORU!$A$2:$D$20,4,FALSE)')
        wh.cell(i, 15, f'=J{i}/3600/(PI()*(N{i}/1000)^2/4)')
        wh.cell(i, 16, f'=23200*d_gaz*J{i}^1.82*N{i}^(-4.82)')
        wh.cell(i, 17, f'=P{i}*K{i}')
        wh.cell(i, 18, h.xi)
        wh.cell(i, 19, f'=0.00397*R{i}*O{i}^2')
        wh.cell(i, 20, f'=-0.049*L{i}')
        wh.cell(i, 21, f'=Q{i}+S{i}+T{i}')
        ust = by_b.get(h.a)
        wh.cell(i, 22, f'=U{i}' + (f'+V{satir[ust.no]}' if ust else ''))
        wh.cell(i, 23, f'=IF(O{i}<=vmax,"UYGUN","HIZ AŞIMI")')
        for c, fmt in ((8, '0.00'), (9, '0.000'), (10, '0.00'), (14, '0.0'), (15, '0.00'), (16, '0.0000'),
                       (17, '0.000'), (19, '0.000'), (20, '0.000'), (21, '0.000'), (22, '0.000')):
            wh.cell(i, c).number_format = fmt
    son = 3 + len(ag.hatlar)
    _ken(wh, 3, son, len(cols))
    wh.conditional_formatting.add(f'W4:W{son}', CellIsRule(operator='equal', formula=['"HIZ AŞIMI"'], fill=KIR))
    wh.conditional_formatting.add(f'W4:W{son}', CellIsRule(operator='equal', formula=['"UYGUN"'], fill=YES))
    for j, w in enumerate((5, 5, 5, 46, 13, 18, 6, 9, 6, 9, 7, 7, 7, 7, 7, 9, 9, 6, 9, 9, 9, 11, 12), 1):
        wh.column_dimensions[CL(j)].width = w
    wh.freeze_panes = 'E4'
    # ------------------------------------------------------------------ KONTROL
    wk = wb.create_sheet('KONTROL')
    wk['A1'] = 'KRİTİK DEVRE KONTROLÜ (sayaç kaybı hariç) - sınırlar ' + AKSA
    wk['A1'].font = Font(bold=True, size=12)
    _baslik(wk, 3, ['Bölüm', 'Cihaz', 'Servis kutusu → sayaç vanası hatları', 'ΔP (mbar)', 'Sınır',
                    'Sayaç → cihaz hatları', 'ΔP (mbar)', 'Sınır', 'Toplam (mbar)', 'Sınır', 'Durum'])
    r = 4
    for d in ag.daireler:
        on = ag.yol_sayac_oncesi(d)
        for hc in d['h_cihaz']:
            so = ag.yol_sayac_sonrasi(d, hc)
            wk.cell(r, 1, d['no'])
            wk.cell(r, 2, hc.cihaz['ad'])
            wk.cell(r, 3, ', '.join(str(h.no) for h in on))
            wk.cell(r, 4, '=' + '+'.join(f'CAP_HESABI!U{satir[h.no]}' for h in on))
            wk.cell(r, 5, '=lim_on')
            wk.cell(r, 6, ', '.join(str(h.no) for h in so))
            wk.cell(r, 7, '=' + '+'.join(f'CAP_HESABI!U{satir[h.no]}' for h in so))
            wk.cell(r, 8, '=lim_so')
            wk.cell(r, 9, f'=D{r}+G{r}')
            wk.cell(r, 10, '=lim_top')
            wk.cell(r, 11, f'=IF(AND(D{r}<=E{r},G{r}<=H{r},I{r}<=J{r}),"UYGUN","UYGUN DEĞİL")')
            for c in (4, 7, 9):
                wk.cell(r, c).number_format = '0.000'
            r += 1
    _ken(wk, 3, r - 1, 11)
    wk.conditional_formatting.add(f'K4:K{r-1}', CellIsRule(operator='equal', formula=['"UYGUN"'], fill=YES))
    wk.conditional_formatting.add(f'K4:K{r-1}', CellIsRule(operator='notEqual', formula=['"UYGUN"'], fill=KIR))
    for c, fx in ((4, 'E'), (7, 'H'), (9, 'J')):
        L_ = CL(c)
        wk.conditional_formatting.add(f'{L_}4:{L_}{r-1}', FormulaRule(formula=[f'{L_}4>{fx}4'], fill=KIR))
    for j, w in enumerate((8, 9, 34, 10, 7, 14, 10, 7, 12, 7, 14), 1):
        wk.column_dimensions[CL(j)].width = w
    # ------------------------------------------------------------------ HAVALANDIRMA
    wv = wb.create_sheet('HAVALANDIRMA')
    wv['A1'] = 'CİHAZ MAHALLİ HAVALANDIRMA KONTROLÜ - ' + AKSA
    wv['A1'].font = Font(bold=True, size=12)
    _baslik(wv, 3, ['Bölüm', 'Cihaz', 'Tip', 'Kapasite (kW)', 'Mahal hacmi (m³)', 'Gerekli hacim (m³)',
                    'Hacim kontrolü', 'Menfez / not'])
    r = 4
    V = P['daire_tipi'].get('mutfak_hacmi_m3', 0)
    for d in ag.daireler:
        for c in d['cihazlar']:
            wv.cell(r, 1, d['no'])
            wv.cell(r, 2, c['ad'])
            wv.cell(r, 3, c['tip'])
            wv.cell(r, 4, '=' + cih_ref[(d['no'], c['ad'])].replace('$G$', '$E$'))
            if c['tip'] == 'A':
                wv.cell(r, 5, V)
                wv.cell(r, 6, f'=MAX(12,D{r}*1)')
                wv.cell(r, 7, f'=IF(E{r}>=F{r},"UYGUN","YETERSİZ")')
                wv.cell(r, 8, 'Sürekli açık ≥150 cm² menfez; yatak odası/banyo/WC/merdivende kurulamaz.')
            elif c['tip'] == 'B':
                wv.cell(r, 5, V)
                wv.cell(r, 6, f'=MAX(8,D{r}*1)')
                wv.cell(r, 7, f'=IF(E{r}>=F{r},"UYGUN","KOMŞU MAHAL 2x150 cm²")')
                wv.cell(r, 8, 'Mahalde ≥150 cm² sürekli açık menfez; CO dedektörü.')
            else:
                wv.cell(r, 7, 'Hermetik')
                wv.cell(r, 8, 'Mahalde dış atmosfere açılan 150 cm² havalandırma; baca doğrudan dış atmosfere.')
            r += 1
    _ken(wv, 3, r - 1, 8)
    wv.conditional_formatting.add(f'G4:G{r-1}', CellIsRule(operator='equal', formula=['"UYGUN"'], fill=YES))
    for j, w in enumerate((8, 9, 5, 12, 14, 15, 18, 70), 1):
        wv.column_dimensions[CL(j)].width = w
    wv.cell(r + 1, 1, 'Mahal hacimleri DENEME değeridir; mimari projeden alınacaktır.' if P.get('deneme') else '')
    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
    wb.save(yol)
    return satir
