import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'YAĞMUR SUYU SARNICI'
b = Side(style='thin'); B = Border(left=b, right=b, top=b, bottom=b)
ws['A1'] = 'ROTAŞ MÜHENDİSLİK - 318-319 PARSEL GÜNGÖREN A-B BLOK'; ws['A1'].font = Font(bold=True, size=13)
ws['A2'] = 'YAĞMUR SUYU SARNICI (YALNIZ BAHÇE SULAMA) - DEPO, FİLTRE, HİDROFOR, TAŞMA HESABI  (PAİY m.57/7-a, TS EN 16941-1, TS EN 12056-3)'
ws['A2'].font = Font(bold=True, size=10)
hdr = ['NO', 'KALEM', 'SEMBOL', 'DEĞER', 'BİRİM', 'AÇIKLAMA']
for i, h in enumerate(hdr):
    c = ws.cell(4, i + 1, h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD'); c.border = B
rows = [
 ('Yıllık yağış (İstanbul, MGM aylık ort. toplamı)', 'h', 659.4, 'mm', 'MGM 1950-2025'),
 ('A Blok çatı iz düşümü (saçak dahil)', 'A_A', 365.0, 'm²', 'çatı planından'),
 ('B Blok çatı iz düşümü (saçak dahil)', 'A_B', 337.7, 'm²', 'çatı planından'),
 ('Toplam toplama alanı', 'A', '=D6+D7', 'm²', ''),
 ('Çatı verim katsayısı (kiremit)', 'e', 0.80, '-', 'TS EN 16941-1'),
 ('Filtre hidrolik verimi', 'η', 0.90, '-', 'F1 girdap filtre'),
 ('Yıllık toplanabilir su E = A·h·e·η', 'E', '=D8*D5/1000*D9*D10', 'm³/yıl', ''),
 ('Asgari depo hacmi Vmin = %6·E', 'Vmin', '=0.06*D11', 'm³', 'PAİY m.57/7-a'),
 ('Sarnıç iç ölçü (mimari: yağmur hasadı deposu) - en', 'a', 3.74, 'm', ''),
 ('Sarnıç iç ölçü - boy', 'b', 5.05, 'm', ''),
 ('Faydalı su yüksekliği', 'hs', 1.50, 'm', ''),
 ('Sarnıç faydalı hacmi', 'V', '=D13*D14*D15', 'm³', ''),
 ('Kontrol V ≥ Vmin', '', '=IF(D16>=D12,"UYGUN","YETERSİZ")', '', ''),
 ('Sulanan bahçe alanı', 'Ab', 181.6, 'm²', 'vaziyet planından'),
 ('Yıllık sulama suyu (60 l/m²·yıl)', 'Ds', '=D18*60/1000', 'm³/yıl', 'DIN 1989-1'),
 ('Kontrol Ds ≤ E', '', '=IF(D19<=D11,"UYGUN","ŞEBEKE TAKVİYESİ ARTAR")', '', ''),
 ('Yaz günlük sulama (5 l/m²·gün)', 'qg', '=D18*5/1000', 'm³/gün', ''),
 ('Depo karşılama süresi', 't', '=D16/D21', 'gün', ''),
 ('Tasarım yağış debisi Qr = A·r·C (r=0,030 l/s·m², C=1)', 'Qr', '=D8*0.03', 'l/s', 'TS EN 12056-3'),
 ('Taşma hattı Ø200 PVC %1 kapasitesi (Manning n=0,010, di=190 mm)', 'Qt', '=PI()*0.19^2/4*(1/0.01)*(0.19/4)^(2/3)*0.01^0.5*1000', 'l/s', 'cazibe, pompa yok'),
 ('Kontrol Qt ≥ Qr', '', '=IF(D24>=D23,"UYGUN","ÇAP ARTIR")', '', 'taşma → yağmur suyu kanalı'),
 ('Hidrofor debisi (1 zon, 6 rotor x 0,25 m³/h)', 'Q', 1.5, 'm³/h', ''),
 ('Geometrik yükseklik', 'Hg', 5.3, 'mSS', '2. bodrum min. su seviyesi → bahçe'),
 ('Sürtünme + filtre + vana kayıpları', 'Hk', 5.7, 'mSS', ''),
 ('Çıkış basıncı (rotor sprinkler)', 'Hç', 25.0, 'mSS', ''),
 ('Basma yüksekliği', 'H', '=D27+D28+D29', 'mSS', ''),
 ('Pompa verimi', 'ηp', 0.40, '-', ''),
 ('Pompa gücü P = ρ·g·Q·H/η', 'P', '=1000*9.81*D26/3600*D30/D31/1000', 'kW', 'seçim 0,55 kW (1 asıl + 1 yedek)'),
 ('Devreye girme basıncı', 'Pd', 3.6, 'bar', ''),
 ('Devreden çıkma basıncı', 'Pk', 5.0, 'bar', ''),
 ('Saatlik devreye girme', 'Z', 20, '1/h', ''),
 ('Membranlı tank V = 1000·Q/(4Z)·(Pk+1)/(Pk−Pd)', 'Vt', '=1000*D26/(4*D35)*(D34+1)/(D34-D33)', 'L', 'seçim 100 L, ön basınç 3,4 bar'),
 ('Kartuş filtre (hidrofor çıkışı)', '', 100, 'µm', 'yıkanabilir'),
 ('F1 girdap filtre', '', 200, 'DN', '0,35 mm elek, çatı ≥ 750 m², kirli su taşma hattına'),
]
r = 5
for i, (k, s, v, u, a) in enumerate(rows, 1):
    vals = [i, k, s, v, u, a]
    for j, x in enumerate(vals):
        c = ws.cell(r, j + 1, x); c.border = B
        if j == 3 and isinstance(x, (int, float)) or (j == 3 and isinstance(x, str) and x.startswith('=') and 'IF' not in x):
            c.number_format = '0.00'
    r += 1
ws.cell(r + 1, 1, 'NOT: Sarnıç taşması depo içindeki taşma sifonundan, F1 kirli suyu ile birlikte cazibeyle YSR rögarı üzerinden yağmur suyu kanalına gider; taşma pompası yoktur.')
ws.cell(r + 2, 1, 'NOT: Yağmur suyu yalnızca bahçe sulamasında kullanılır, içilemez; pis su şebekesine bağlanmaz. Şebeke takviyesi TS EN 1717 AB hava boşluğu ile yapılır.')
for col, w in zip('ABCDEF', (5, 62, 8, 12, 9, 40)):
    ws.column_dimensions[col].width = w
ws.print_area = 'A1:F%d' % (r + 2)
ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save('out/318-319_GUNGOREN_YAGMUR_SUYU_SARNICI_HESAP.xlsx')
