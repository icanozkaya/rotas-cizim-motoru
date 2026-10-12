import sys, openpyxl
sys.path.insert(0,'/home/claude/nb/p4'); sys.path.insert(0,'/home/claude/nb'); sys.path.insert(0,'/home/claude/nb/p2')
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import importlib.util
spec = importlib.util.spec_from_file_location('kdm', '/home/claude/nb/p4/kd.py'); kdm = importlib.util.module_from_spec(spec); spec.loader.exec_module(kdm)
b = Side(style='thin'); B = Border(left=b, right=b, top=b, bottom=b)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'EKİPMAN LİSTESİ'
ws['A1'] = 'ROTAŞ MÜHENDİSLİK - KAZAN DAİRESİ EKİPMAN LİSTESİ'; ws['A1'].font = Font(bold=True, size=13)
for i, h in enumerate(['NO', 'ADI', 'ADEDİ', 'ÖZELLİĞİ']):
    c = ws.cell(3, i + 1, h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD'); c.border = B
for r, row in enumerate(kdm.EKIPMAN, 4):
    for j, v in enumerate(row):
        c = ws.cell(r, j + 1, v); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='center')
for col, w in zip('ABCD', (6, 34, 8, 70)): ws.column_dimensions[col].width = w
ws.print_area = 'A1:D%d' % (3 + len(kdm.EKIPMAN)); ws.sheet_properties.pageSetUpPr.fitToPage = True
h = wb.create_sheet('HESAP KONTROL')
h['A1'] = 'KAZAN DAİRESİ HESAP KONTROLLERİ'; h['A1'].font = Font(bold=True, size=13)
rows = [
 ('Kazan adedi', 6, 'adet'), ('Kazan kapasitesi', 115, 'kW'), ('Toplam kazan kapasitesi', '=B3*B4', 'kW'),
 ('Eşanjör kapasitesi', 700, 'kW'), ('Kontrol', '=IF(B5>=0.98*B6,"UYGUN","YETERSİZ")', ''),
 ('Primer gidiş / dönüş', '70/50', '°C'), ('ΔT', 20, 'K'), ('Primer debi Q = P/(1,163·ΔT)', '=B6/(1.163*B9)', 'm³/h'),
 ('Seçilen sekonder pompa', 30, 'm³/h'), ('Kontrol', '=IF(B11>=0.99*B10,"UYGUN","ARTIR")', ''),
 ('3" boru iç çapı', 80.8, 'mm'), ('Kollektör hızı', '=B11/3600/(PI()*(B13/1000)^2/4)', 'm/s'),
 ('Alt havalandırma alanı (6 cm²/kW)', '=B5*6', 'cm²'), ('Seçilen alt menfez 70x60', '=70*60', 'cm²'),
 ('Üst havalandırma alanı (3 cm²/kW)', '=B5*3', 'cm²'), ('Seçilen üst menfez 50x45', '=50*45', 'cm²'),
 ('Akümülasyon hacmi', '=7*1000', 'L'), ('Isıtma süresi 10→60 °C', '=B19*1.163*50/(B6*1000)*60', 'dk'),
 ('Hidrofor debisi', 29, 'm³/h'), ('Su deposu', 30, 'ton'), ('Depo / hidrofor debisi', '=B22/B21', 'saat'),
]
for i, (a, v, u) in enumerate(rows, 3):
    for j, x in enumerate((a, v, u)):
        c = h.cell(i, j + 1, x); c.border = B
        if j == 1: c.number_format = '0.00'
for col, w in zip('ABC', (40, 14, 8)): h.column_dimensions[col].width = w
h.print_area = 'A1:C%d' % (2 + len(rows)); h.sheet_properties.pageSetUpPr.fitToPage = True
wb.save('out/KAZAN_DAIRESI_EKIPMAN_HESAP.xlsx')
