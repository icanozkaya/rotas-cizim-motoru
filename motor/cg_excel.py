"""Cafer Göktaş: ISI KAYBI.xls ve YERDEN ISITMA HESABI.xls şablonlarını doldur."""
import pickle, sys, math, copy
sys.path.insert(0, '/home/claude/nb')
import openpyxl
import cg_heat

OUT = '/tmp/claude-0/xl'
R, fr, HEAT = pickle.load(open('/tmp/claude-0/cg_recs.pkl', 'rb'))
FLOORS = [('Z', 'ZEMİN KAT'), ('1', '1. KAT'), ('2', '2. KAT')]
ORDER = ['ODA', 'ODA2', 'BANYO + WC', 'SALON VE MUTFAK', 'ANTRE']
DISP = {'ODA': 'ODA', 'ODA2': 'ODA', 'BANYO + WC': 'BANYO + WC', 'SALON VE MUTFAK': 'SALON VE MUTFAK', 'ANTRE': 'ANTRE'}


def zl(key, nm):
    z = cg_heat.ZNO[nm]
    return ('Z%02d' % z) if key == 'Z' else ('%s%02d' % (key, z))


def zl_sp(key, nm):
    z = cg_heat.ZNO[nm]
    return ('Z %02d' % z) if key == 'Z' else ('%s %02d' % (key, z))


# ================================================================== ISI KAYBI
wb = openpyxl.load_workbook(OUT + '/isi.xlsx')
ws = wb['KALORİFER  KAPAK']
ws['C4'] = 'CAFER GÖKTAŞ'; ws['C5'] = 'KONUT'; ws['C6'] = 3
ws['B9'] = 'SAKARYA'; ws['B10'] = 'ADAPAZARI'; ws['B11'] = 'PAPUÇÇULAR'; ws['B12'] = ''
ws['B13'] = 'G24-b-23-c-2-b'; ws['B14'] = 405; ws['B15'] = 381
ws['B18'] = 'BAYRAM GENÇ'; ws['B19'] = 'MAKİNA MÜH.'

ws = wb['BİNAYA İLİŞKİN BELGELER']
ws['A4'] = '•Bina zemin + 2 normal kat olmak üzere 3 katlıdır.'
ws['A5'] = '•Binada 3 daire ısıtma yapılıyor (her katta 1 daire).'
ws['A6'] = '•Bina doğu cephesinde “bitişik nizamda”, “normal bölgede” serbest olarak alındı.'
ws['A8'] = '•Mahal iç sıcaklıkları 20 ve 24 derecedir; merdiven holü 10, asansör boşluğu 5 derece alınmıştır.'
ws['A9'] = '•Sakarya (Adapazarı) için dış sıcaklık -3 derece alınmıştır.'
ws['A10'] = '•Çatı arası sıcaklığı 4 derece, zemin altı (toprak) sıcaklığı 9 derecedir.'
ws['A21'] = ' Kombi - kollektör arası çok katmanlı/PPRC, döşeme devreleri PEX 16x2.0 borudur.'
ws['A22'] = 'Binada 3 adet duvar tipi, yoğuşmalı, hermetik, 24 kW’lık kombi ve yerden ısıtma sistemi kullanılmıştır.'
ws['A30'] = 'DD       DIŞ DUVAR          =    0,493  W/m2 K'
ws['A34'] = 'DK        DAİRE KAPISI         =    3,500  W/m2 K'
ws['A38'] = 'DÖ        DÖŞEME                =    0,231  W/m2 K'
ws['A42'] = 'İD        İÇ DUVAR              =    0,402 W/m2 K'

ws = wb['ISI KAYBI HESABI']
heads = [c.row for c in ws['A'] if c.value == 'ISI KAYBI HESAP CETVELİ']
_body = set()
for h in heads[:12]:
    _body |= set(range(h + 15, h + 49))
for mr in list(ws.merged_cells.ranges):
    if mr.min_row in _body:
        ws.unmerge_cells(str(mr))
# eski verileri temizle (gövde satırları)
for h in heads[:12]:
    for r in range(h + 15, h + 49):
        for col in range(1, 18):
            cell = ws.cell(r, col)
            if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                cell.value = None

blocks = []
for key, fname in FLOORS:
    c = cg_heat.calc(key)
    for nm in ORDER:
        blocks.append((key, fname, nm, c[nm]))
TOTALS = {}
page, row = 0, heads[0] + 15
used = 0
for key, fname, nm, d in blocks:
    need = len(d['rows']) + 5
    if row + need > heads[page] + 48:
        page += 1
        row = heads[page] + 15
    h = heads[page]
    ws.cell(h + 1, 17).value = fname
    ws.cell(row, 4).value = '%s %s (%d°C)' % (zl_sp(key, nm), DISP[nm], d['ti'])
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=11)
    r0 = row + 1
    cp_rows = []
    for i, x in enumerate(d['rows']):
        r = r0 + i
        ws.cell(r, 1).value = x['mark']
        if x['l'] is not None:
            ws.cell(r, 4).value = round(x['l'], 2)
            ws.cell(r, 5).value = round(x['h'], 2)
            ws.cell(r, 6).value = '=SUM(D%d*E%d)' % (r, r)
        else:
            ws.cell(r, 6).value = round(x['F'], 2)
        ws.cell(r, 7).value = x['n']
        if x['mark'] == 'ÇP':
            cp_rows.append(r)
        if x['minus']:
            src = [r0 + j for j, y in enumerate(d['rows'][:i]) if y['mark'] == x.get('ref')]
            ws.cell(r, 8).value = '=SUM(%s)' % ','.join('I%d' % s for s in src) if src else round(x['minus'], 2)
        else:
            ws.cell(r, 8).value = 0
        ws.cell(r, 9).value = '=SUM(F%d*G%d-H%d)' % (r, r, r)
        ws.cell(r, 10).value = x['U']
        ws.cell(r, 11).value = x['dt']
        ws.cell(r, 17).value = '=SUM(I%d*J%d*K%d)' % (r, r, r)
    rs = r0 + len(d['rows'])
    ws.cell(rs, 17).value = '=SUM(Q%d:Q%d)' % (r0, rs - 1)
    ra = rs + 1
    ws.cell(ra, 4).value = 'Q s = ∑ a × L × R × H × ∆t × Ze'
    ws.merge_cells(start_row=ra, start_column=4, end_row=ra, end_column=11)
    ws.cell(ra, 13).value = cg_heat.ZD; ws.cell(ra, 14).value = cg_heat.ZW; ws.cell(ra, 15).value = cg_heat.ZH
    ws.cell(ra, 16).value = '=1+(M%d+N%d+O%d)/100' % (ra, ra, ra)
    ws.cell(ra, 17).value = '=SUM(Q%d*P%d)' % (rs, ra)
    rq = ra + 1
    ws.cell(rq, 5).value = 'Qs= '
    ws.cell(rq, 6).value = ('=' + '+'.join('(D%d+E%d*G%d)' % (r, r, r) for r in cp_rows)) if cp_rows else 0
    ws.cell(rq, 7).value = 2; ws.cell(rq, 8).value = 0.97; ws.cell(rq, 9).value = d['dto']
    ws.cell(rq, 10).value = 0.9; ws.cell(rq, 11).value = 1
    ws.cell(rq, 17).value = '=SUM(F%d*G%d*H%d*I%d*J%d*K%d)' % (rq, rq, rq, rq, rq, rq)
    rt = rq + 1
    ws.cell(rt, 17).value = '=SUM(Q%d:Q%d)' % (ra, rq)
    TOTALS[(key, nm)] = ("'ISI KAYBI HESABI'!Q%d" % rt, d['Q'])
    row = rt + 2

# radyatör (donanım) seçimi -> yerden ısıtma mahal listesi
ws = wb['RADYATÖR SEÇİMİ']
for r in range(15, 45):
    for col in range(1, 22):
        cell = ws.cell(r, col)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
r = 15
for k, (key, fname) in enumerate(FLOORS):
    ws.cell(r, 1).value = 'D%d' % (k + 1); ws.cell(r, 2).value = fname
    r += 1
    first = r
    for nm in ORDER:
        ref, q = TOTALS[(key, nm)]
        ws.cell(r, 1).value = zl(key, nm); ws.cell(r, 2).value = DISP[nm]
        ws.cell(r, 3).value = 24 if 'BANYO' in nm else 20
        ws.cell(r, 5).value = '=' + ref
        ws.cell(r, 9).value = 'YERDEN ISITMA (PEX 16x2.0 / 15 cm)'
        r += 1
    ws.cell(r, 5).value = '=SUM(E%d:E%d)' % (first, r - 1)
    ws.cell(r, 2).value = 'DAİRE %d TOPLAM' % (k + 1)
    r += 2

# boru çapı: kombi-kollektör hatları
ws = wb['BORU ÇAPI HESABI']
for rr in range(13, 23):
    for col in range(1, 19):
        cell = ws.cell(rr, col)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
for k, (key, fname) in enumerate(FLOORS):
    res, tot = HEAT[k]
    rr = 13 + k
    m = tot / (4186 * 10)                      # kg/s, ΔT=10K
    din = 0.016                                 # DN20 çok katmanlı iç çap
    v = m / 1000 / (math.pi * din * din / 4)
    Rpa = 0.0  # Darcy yaklaşık
    lam = 0.03
    Rpa = lam / din * 1000 * v * v / 2
    L = 9.5
    ws.cell(rr, 1).value = 'D%d kombi-kollektör' % (k + 1)
    ws.cell(rr, 2).value = round(tot, 1)
    ws.cell(rr, 4).value = L
    ws.cell(rr, 5).value = '3/4"'
    ws.cell(rr, 12).value = round(v, 2)
    ws.cell(rr, 13).value = round(Rpa, 1)
    ws.cell(rr, 14).value = round(Rpa * L, 1)
    ws.cell(rr, 15).value = 8.0
    ws.cell(rr, 16).value = round(8.0 * 1000 * v * v / 2, 1)
ws.cell(19, 1).value = 'En uzun devre: PEX 16x2.0, ~48 m, 588 W; R≈90 Pa/m -> ~4.3 kPa + kollektör/vana 10 kPa'
ws = wb['TO DEĞERLERİ HESABI']
for rr in range(19, 40):
    for col in range(1, 23):
        cell = ws.cell(rr, col)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
ws.cell(33, 3).value = 'H = devre (≈4.3 kPa) + kollektör, vanalar ve kombi iç direnci (≈10 kPa) ≈ 15 kPa = 1.5 mSS'
ws.cell(36, 3).value = 'Kombi sirkülasyon pompası (yaklaşık 5-6 mSS) yerden ısıtma devreleri için yeterlidir.'
ws.cell(39, 3).value = 'Her daire kombisi ayrı devre olduğundan ortak sirkülasyon pompası gerekmez.'
wb.save(OUT + '/CAFER_GOKTAS_ISI_KAYBI.xlsx')

# ================================================================== YERDEN ISITMA
DIMS = {  # en, boy (m), kenar bölge boyu (dış duvar, m)
    'ODA': (2.46, 4.47, 4.47), 'ODA2': (2.13, 2.19, 2.13), 'BANYO + WC': (1.78, 2.25, 0.0),
    'SALON VE MUTFAK': (3.06, 6.21, 6.21), 'ANTRE': (1.10, 4.20, 0.0)}
wb = openpyxl.load_workbook(OUT + '/yi.xlsx')
ws = wb.active
_mer = [str(m) for m in ws.merged_cells.ranges]
for m in _mer:
    ws.unmerge_cells(m)
tmpl = {c: copy.copy(ws.cell(8, c)._style) for c in range(1, 18)}
tmpl2 = {c: copy.copy(ws.cell(9, c)._style) for c in range(1, 18)}
hdr = {r: {c: (ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style)) for c in range(1, 18)} for r in (5, 6, 7)}
for r in range(1, 200):
    for c in range(1, 18):
        cell = ws.cell(r, c)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = None
ws['G2'] = 'CAFER GÖKTAŞ'
r = 4
for k, (key, fname) in enumerate(FLOORS):
    res, tot = HEAT[k]
    c = cg_heat.calc(key)
    ws.cell(r, 1).value = 'DAİRE %d (%s)' % (k + 1, fname)
    r += 1
    for hr in (5, 6, 7):
        for col in range(1, 18):
            v, st = hdr[hr][col]
            ws.cell(r, col).value = v
            ws.cell(r, col)._style = copy.copy(st)
        r += 1
    first = r
    for nm in ORDER:
        lens = [d['L'] for d in res if d['room'] == nm]
        if not lens:
            lens = [0.0]
        nrow = max(2, len(lens))
        en, boy, kb = DIMS[nm]
        kcal = round(c[nm]['Q'] * 0.86, 0)
        r1 = r
        for j in range(nrow):
            rr = r + j
            for col in range(1, 18):
                ws.cell(rr, col)._style = copy.copy(tmpl[col] if j == 0 else tmpl2[col])
            ws.cell(rr, 12).value = 'M15'
            ws.cell(rr, 13).value = '=IF(L%d="M10",8,IF(L%d="M15",6,IF(L%d="M20",4.5,IF(L%d="M25",3.5,IF(L%d="M30",3,IF(L%d="M35",2.5))))))' % ((rr,) * 6)
            ws.cell(rr, 14).value = round(lens[j], 2) if j < len(lens) else 0
        ws.cell(r1, 1).value = zl(key, nm)
        ws.cell(r1, 2).value = DISP[nm] + (' (geçiş boruları)' if nm == 'ANTRE' else '')
        ws.cell(r1, 3).value = en; ws.cell(r1, 4).value = boy; ws.cell(r1, 5).value = 3.0
        ws.cell(r1, 6).value = kcal
        # kenar bölge
        ws.cell(r1, 7).value = 0.5 if kb else 0
        ws.cell(r1, 8).value = kb
        ws.cell(r1, 9).value = '=G%d*H%d' % (r1, r1)
        ws.cell(r1, 10).value = '=I%d*K%d' % (r1, r1)
        ws.cell(r1, 11).value = 150
        # iç bölge
        r2 = r1 + 1
        ws.cell(r2, 7).value = round(en - (0.5 if kb else 0), 2)
        ws.cell(r2, 8).value = boy
        ws.cell(r2, 9).value = '=G%d*H%d' % (r2, r2)
        ws.cell(r2, 10).value = '=F%d-J%d' % (r1, r1)
        ws.cell(r2, 11).value = '=IF(I%d>0,J%d/I%d,0)' % (r2, r2, r2)
        ws.cell(r1, 15).value = '=SUM(N%d:N%d)' % (r1, r1 + nrow - 1)
        ws.cell(r1, 16).value = '=O%d/80' % r1
        ws.cell(r1, 17).value = len([d for d in res if d['room'] == nm])
        r += nrow
    ws.cell(r, 14).value = '=SUM(N%d:N%d)' % (first, r - 1)
    ws.cell(r, 13).value = 'TOPLAM'
    r += 3
wb.save(OUT + '/CAFER_GOKTAS_YERDEN_ISITMA.xlsx')
print('ok', TOTALS[('Z', 'SALON VE MUTFAK')])
