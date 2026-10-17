# -*- coding: utf-8 -*-
"""Adapeker — ISI KAYBI + YERDEN ISITMA hesap Excel'i (kullanıcının şablonu sablon/ornek_yi.xlsx; ısı kaybı sayfaları
excel_yi.isi_kaybi_sayfalari ile ZEMİN / 1 NK düzeninde, formüllü). Girdi: proje/adapeker/yerden_model.json"""
import os, sys, json, shutil, datetime
BURASI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BURASI)
from ortak_ad import KOK, CIKIS, PROJE
sys.path.insert(0, os.path.join(KOK, 'motor', 'yerden2'))
import excel_yi as EX  # noqa
import isi_kaybi as IK  # noqa
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

SABLON = os.path.join(KOK, 'sablon', 'ornek_yi.xlsx')
B = Border(*(Side(style='thin'),) * 4)


def main():
    J = json.load(open(os.path.join(CIKIS, 'yerden_model.json')))
    yol = os.path.join(CIKIS, 'ADAPEKER_ISI_KAYBI_YERDEN_ISITMA_HESABI.xlsx')
    shutil.copyfile(SABLON, yol)
    wb = openpyxl.load_workbook(yol)
    tpl = openpyxl.load_workbook(SABLON)
    # ---------------- kapak
    ws = wb['kapak']
    ws['A2'] = ('ROTAŞ MÜHENDİSLİK.PROJE.MÜŞAVİRLİK.\nSemerciler Mah. Sait Faik Sk. No:22/C Z01 Adapazarı/Sakarya\n'
                'rotasmuhendislik@gmail.com - 0264 272 00 49\n')
    ws['A10'] = 'Teknik Ressam'
    ws['B10'] = 'CAN ÖZKAYA'
    ws['D10'] = None
    ws['F10'] = None
    ws['A11'] = 'ISI KAYBI VE YERDEN ISITMA TESİSATI PROJE RAPORU'
    ws['C31'] = PROJE['sahibi']
    ws['A36'], ws['B36'], ws['D36'], ws['E36'] = PROJE['il'], PROJE['ilce'], PROJE['mahalle'], '-'
    ws['G36'], ws['H36'], ws['I36'] = PROJE['pafta'], PROJE['ada'], PROJE['parsel']
    ws['A41'] = 'BETONARME'
    ws['B41'] = PROJE['yapi_sinifi']
    ws['C41'] = PROJE['bb']
    ws['D41'] = 'ZEMİN + 2'
    ws['F41'] = None
    ws['G41'] = datetime.datetime(2026, 10, 18)
    ws['H41'] = 'MESKEN'
    # ---------------- teknik rapor
    ws = wb['TEKNİK RAPOR']
    ws['A1'] = 'ISI KAYBI VE YERDEN ISITMA TESİSATI PROJESİ TEKNİK RAPORU'
    ws['E10'] = ': Mesken (Zemin + 2 kat, 6 bağımsız bölüm, kat başına 2 daire)'
    ws['E11'] = ': Duvar tipi yoğuşmalı kombi (her dairede 1 adet, 24 kW, balkonda - mimari yerinde)'
    ws['E13'] = ': Zeminden ısıtma, PEX 16x2,0 (oksijen bariyerli), boru aralığı 15 cm, 45/35 °C'
    ws['E16'] = ': Zemin + 2 normal kat, kat yüksekliği 3,00 m (mimari kesit: +0,50 / +3,50 / +6,50 / çatı +9,50)'
    ws['E17'] = ': PEX 16 x 2,0'
    ws['E18'] = ': 45 - 35 °C'
    ws['B21'] = 'Salon'; ws['F21'] = ': 20°C'
    ws['B22'] = 'Oda / Yatak Odası'; ws['F22'] = ': 20°C'
    ws['B23'] = 'Mutfak'; ws['F23'] = ': 20°C'
    ws['B24'] = 'Banyo - WC'; ws['F24'] = ': 24°C'
    ws['B25'] = 'Hol (daire içi)'; ws['F25'] = ': 15°C (yerden ısıtma yok, ayaklar geçer)'
    for c in ('F26', 'F27', 'F28'):
        ws[c] = None
    ws['B31'] = 'Isıtılmayan'; ws['F31'] = ': balkon, çamaşır odası (1,60 m²), merdiven, kat holü, asansör'
    for rr in range(33, 70):
        for cc in range(1, 15):
            ws.cell(rr, cc).value = None
    U = IK.u_sec(3)
    sat = [('1.3.', 'Isı Yalıtımı (TS 825:2024, İZODER il sayfası - SAKARYA 3. BÖLGE):'),
           ('', '   Dış duvar U = %.2f W/m²K, çatı arası tavanı U = %.2f, toprağa oturan döşeme U = %.2f, pencere / balkon kapısı U = %.2f.'
            % (U['DD'], U['Ta'], U['Dö'], U['ÇP'])),
           ('', '   Bu değerler TS 825:2024 3. bölge tavsiye edilen (aşılmayacak) sınır değerlerdir; yalıtım detayları bunları aşmayacak.'),
           ('', '   Dış hava -3 °C (Adapazarı, TS 2164), toprak 6 °C, merdiven/ortak alan 10 °C, hol 15 °C; kat yüksekliği 3,00 m.'),
           ('1.4.', 'Yerden Isıtma Esası:'),
           ('', '   Boru metresi başına 35 kcal/h (devre kapasitesi = 35 kcal/h × L); devre boyu en çok 80 m, 15 cm aralık.'),
           ('', '   Her mahalde boru kapasitesi ≥ ısı kaybı kontrol edilmiştir (YERDEN ISITMA HESABI sayfası).'),
           ('', '   Hazırlayan: Can ÖZKAYA - Teknik Ressam (ROTAŞ Mühendislik).')]
    for i, (a, b) in enumerate(sat):
        ws['A%d' % (33 + 2 * i)] = a
        ws['B%d' % (33 + 2 * i)] = b
    # ---------------- ısı kaybı sayfaları
    katlar, mahaller = [], []
    for kt in J['model']:
        kat = dict(kat=kt['kat'], ad={'Z': 'ZEMİN KAT', '1': '1. KAT', '2': '2. KAT'}[kt['kat']])
        katlar.append(kat)
        for n in sorted(kt['daireler'], key=int):
            d = kt['daireler'][n]
            for m in d['mahaller']:
                mahaller.append((kat, d['no'], m))
    ref = EX.isi_kaybi_sayfalari(wb, tpl, katlar, mahaller)
    for nm in ('ZEMİN', '1 NK', 'TEKNİK RAPOR (2)', 'PETEK LİST', 'YERDEN ISITMA', 'MMO YERDEN ISITMA', 'YERDEN ISITMA- DEVAMI'):
        if nm in wb.sheetnames:
            del wb[nm]
    for ws in wb.worksheets:
        if ws.title.startswith('ISI KAYBI'):
            ws['H2'] = PROJE['sahibi']
    # ---------------- yerden ısıtma hesabı (mahal) + devre listesi
    ws = wb.create_sheet('YERDEN ISITMA HESABI')
    ws['A1'] = 'YERDEN ISITMA HESABI - MAHAL ISI KAYBI / BORU KAPASİTESİ (35 kcal/h · m)'
    ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = PROJE['ad'] + ' - Ada %s Parsel %s - Sakarya 3. bölge' % (PROJE['ada'], PROJE['parsel'])
    bas = ['KAT', 'DAİRE', 'ZON', 'MAHAL', 'Ti °C', 'ALAN m²', 'ISI KAYBI W', 'ISI KAYBI kcal/h', 'BORU L (m)', 'DEVRE',
           'KAPASİTE kcal/h', 'DURUM']
    for j, b in enumerate(bas):
        c = ws.cell(4, j + 1, b)
        c.font = Font(bold=True); c.border = B; c.fill = PatternFill('solid', fgColor='DDE6F0')
        c.alignment = Alignment(wrap_text=True, horizontal='center')
    r = 5
    for row in J['ozet']:
        kat = next(k for k in katlar if k['kat'] == row['kat'])
        v = [kat['ad'], row['daire'], row['zon'], row['ad'], row['sic'], row['alan']]
        for j, x in enumerate(v):
            ws.cell(r, j + 1, x).border = B
        rf = ref.get((row['kat'], row['zon']))
        ws.cell(r, 7, ('=' + rf) if rf else row['Q_W']).border = B
        ws.cell(r, 8, '=ROUND(G%d/1.163,1)' % r).border = B
        ws.cell(r, 9, row['L']).border = B
        ws.cell(r, 10, row['n']).border = B
        ws.cell(r, 11, '=ROUND(35*I%d,0)' % r).border = B
        ws.cell(r, 12, '=IF(K%d>=H%d,"YETERLİ","İLAVE GEREKLİ")' % (r, r)).border = B
        r += 1
    ws.cell(r, 6, 'TOPLAM').font = Font(bold=True)
    ws.cell(r, 7, '=SUM(G5:G%d)' % (r - 1)); ws.cell(r, 8, '=SUM(H5:H%d)' % (r - 1))
    ws.cell(r, 9, '=SUM(I5:I%d)' % (r - 1)); ws.cell(r, 11, '=SUM(K5:K%d)' % (r - 1))
    for col, wdt in zip('ABCDEFGHIJKL', (11, 7, 7, 16, 7, 9, 12, 13, 11, 7, 14, 15)):
        ws.column_dimensions[col].width = wdt
    # devre listesi
    ws2 = wb.create_sheet('DEVRE LİSTESİ')
    ws2['A1'] = 'YERDEN ISITMA DEVRE LİSTESİ (kollektör ağız sırası şemadaki gibi)'
    ws2['A1'].font = Font(bold=True, size=12)
    for j, b in enumerate(['KAT', 'DAİRE', 'ZON', 'MAHAL', 'DEVRE NO', 'L (m)', 'ARALIK cm', 'KAPASİTE kcal/h', 'MAHAL PAYI kcal/h']):
        c = ws2.cell(3, j + 1, b); c.font = Font(bold=True); c.border = B
    r = 4
    for kt in J['model']:
        for n in sorted(kt['daireler'], key=int):
            d = kt['daireler'][n]
            for m in d['mahaller']:
                for dv in m['devreler']:
                    for j, x in enumerate([kt['kat'], d['no'], m['zon'], m['ad'], dv['no'], dv['L_m'], 15, dv['kcal'], dv['q_mahal']]):
                        ws2.cell(r, j + 1, x).border = B
                    r += 1
    wb.save(yol)
    print(yol, wb.sheetnames)


if __name__ == '__main__':
    main()
